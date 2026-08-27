"""Checksum-qualified binding for the official Test 2 diagnostic case 2E1.

The supplied Test 2 workbook contains a candidate transfer sheet and several
reference-program sheets.  This module binds the eight hourly series needed by
the fixed-closed fabric-awning diagnostic to their exact workbook columns.

Only the workbook schema is qualified here.  No candidate values are written,
no reference envelope is inferred, no VE result is accepted, and the dynamic
Test 2A control remains outside this contract.
"""

from __future__ import annotations

import hashlib
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Tuple, Union

from openpyxl import load_workbook

from ..exceptions import ConfigurationError

EXPECTED_TEST2_WORKBOOK_SHA256 = (
    "0d34793b0a193e0fac52918a359ac50809046fe7bb56dbab4412ed42332d4d1e"
)

CANDIDATE_SHEET = "Daten_Testprogramm"
CANDIDATE_DATA_START_ROW = 5
REFERENCE_DATA_START_ROW = 4
HOUR_COUNT = 8760

_IRRADIATION_LABELS = {
    "hourly_incident_solar_irradiance_total_on_window_plane": (
        "Solare Einstrahlung auf die Fensterebene gesamt (W/m^2)"
    ),
    "hourly_incident_solar_irradiance_direct_on_window_plane": (
        "Solare Einstrahlung auf die Fensterebene direkt (W/m^2)"
    ),
    "hourly_incident_solar_irradiance_diffuse_on_window_plane": (
        "Solare Einstrahlung auf die Fensterebene diffus (W/m^2)"
    ),
}

_DIAGNOSTIC_LABELS = {
    "hourly_room_solar_heat_gain_total": "Solarer Wärmeeintrag gesamt (W)",
    "hourly_room_solar_heat_gain_direct": (
        "Direkter solarer Wärmeeintrag Direktstrahlung, (W)"
    ),
    "hourly_room_solar_heat_gain_diffuse": (
        "Direkter solarer Wärmeeintrag Diffusstrahlung, (W)"
    ),
    "hourly_room_solar_heat_gain_secondary": ("Sekundärer solarer Wärmeeintrag (W)"),
    "hourly_transmitted_solar_radiation_excluding_secondary": (
        "Total transmittierte Solarstrahlung\n(W)"
    ),
}

_SERIES_COLUMNS = (
    (
        "hourly_incident_solar_irradiance_total_on_window_plane",
        "Y",
    ),
    (
        "hourly_incident_solar_irradiance_direct_on_window_plane",
        "Z",
    ),
    (
        "hourly_incident_solar_irradiance_diffuse_on_window_plane",
        "AA",
    ),
    ("hourly_room_solar_heat_gain_total", "BD"),
    ("hourly_room_solar_heat_gain_direct", "BE"),
    ("hourly_room_solar_heat_gain_diffuse", "BF"),
    ("hourly_room_solar_heat_gain_secondary", "BG"),
    (
        "hourly_transmitted_solar_radiation_excluding_secondary",
        "BH",
    ),
)


def _sha256(path: Path) -> str:
    """Return the lowercase SHA-256 of one workbook."""

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normal(value: Any) -> str:
    """Normalize only whitespace needed for robust exact-header comparison."""

    return " ".join(str(value or "").replace("\n", " ").split())


@dataclass(frozen=True)
class Test2AWorkbookSeriesBinding:
    """One exact 2E1 quantity-to-column binding."""

    series_id: str
    unit: str
    candidate_column: str
    candidate_header_cell: str
    candidate_data_range: str
    reference_column: str
    reference_header_cell: str
    reference_data_range: str


@dataclass(frozen=True)
class Test2ADiagnosticWorkbookBinding:
    """Verified workbook identity and complete 2E1 transfer schema."""

    workbook_path: Path
    workbook_sha256: str
    candidate_sheet: str
    candidate_case_marker_cell: str
    candidate_timestamp_range: str
    reference_sheet_prefix: str
    reference_case_marker_cell: str
    reference_timestamp_range: str
    hour_count: int
    series: Tuple[Test2AWorkbookSeriesBinding, ...]
    status: str = "OFFICIAL_2E1_SCHEMA_BOUND"

    def to_dict(self) -> Dict[str, Any]:
        """Return a JSON-safe, audit-ready binding."""

        payload = asdict(self)
        payload["workbook_path"] = str(self.workbook_path)
        payload["series"] = [asdict(item) for item in self.series]
        payload["claim_guardrail"] = (
            "This binds the official workbook schema only. Empty candidate "
            "cells, VE optical behaviour, reference-program scatter and "
            "dynamic Test 2A control are not qualified by this binding."
        )
        return payload


def _expected_label(series_id: str) -> str:
    """Return the exact official header for one bound series."""

    if series_id in _IRRADIATION_LABELS:
        return _IRRADIATION_LABELS[series_id]
    return _DIAGNOSTIC_LABELS[series_id]


def _series_unit(series_id: str) -> str:
    """Return the unit declared in the official header."""

    return "W/m2" if series_id in _IRRADIATION_LABELS else "W"


def _assert_cell(
    worksheet: Any,
    cell: str,
    expected: str,
    *,
    context: str,
) -> None:
    """Fail closed when one schema-defining cell differs."""

    actual = worksheet[cell].value
    if _normal(actual) != _normal(expected):
        raise ConfigurationError(
            "{} schema mismatch at {}: expected {!r}, found {!r}".format(
                context, cell, expected, actual
            )
        )


@lru_cache(maxsize=8)
def _load_cached(
    workbook_path: str,
    expected_sha256: str,
) -> Test2ADiagnosticWorkbookBinding:
    """Load and verify one immutable official workbook binding."""

    path = Path(workbook_path)
    if not path.is_file():
        raise ConfigurationError("Official Test 2 workbook is missing: {}".format(path))
    actual_sha = _sha256(path)
    expected = expected_sha256.lower()
    if actual_sha != expected:
        raise ConfigurationError(
            "Official Test 2 workbook SHA-256 mismatch: expected {}, found {}".format(
                expected, actual_sha
            )
        )

    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if CANDIDATE_SHEET not in workbook.sheetnames:
            raise ConfigurationError(
                "Official Test 2 workbook is missing sheet {!r}".format(CANDIDATE_SHEET)
            )
        candidate = workbook[CANDIDATE_SHEET]
        _assert_cell(
            candidate,
            "X3",
            "Date/Time",
            context=CANDIDATE_SHEET,
        )
        _assert_cell(
            candidate,
            "BD2",
            "2 E1",
            context=CANDIDATE_SHEET,
        )
        for series_id, column in _SERIES_COLUMNS:
            _assert_cell(
                candidate,
                "{}3".format(column),
                _expected_label(series_id),
                context=CANDIDATE_SHEET,
            )

        reference_names = tuple(
            name
            for name in workbook.sheetnames
            if name.startswith("Daten_") and name != CANDIDATE_SHEET
        )
        if not reference_names:
            raise ConfigurationError(
                "Official Test 2 workbook contains no reference Daten_* sheet"
            )
        for name in reference_names:
            worksheet = workbook[name]
            _assert_cell(
                worksheet,
                "X2",
                "Date/Time",
                context=name,
            )
            _assert_cell(
                worksheet,
                "BC2",
                "2 E1",
                context=name,
            )
            for series_id, column in _SERIES_COLUMNS:
                _assert_cell(
                    worksheet,
                    "{}2".format(column),
                    _expected_label(series_id),
                    context=name,
                )
    finally:
        workbook.close()

    candidate_end = CANDIDATE_DATA_START_ROW + HOUR_COUNT - 1
    reference_end = REFERENCE_DATA_START_ROW + HOUR_COUNT - 1
    series = tuple(
        Test2AWorkbookSeriesBinding(
            series_id=series_id,
            unit=_series_unit(series_id),
            candidate_column=column,
            candidate_header_cell="{}3".format(column),
            candidate_data_range="{}{}:{}{}".format(
                column,
                CANDIDATE_DATA_START_ROW,
                column,
                candidate_end,
            ),
            reference_column=column,
            reference_header_cell="{}2".format(column),
            reference_data_range="{}{}:{}{}".format(
                column,
                REFERENCE_DATA_START_ROW,
                column,
                reference_end,
            ),
        )
        for series_id, column in _SERIES_COLUMNS
    )
    return Test2ADiagnosticWorkbookBinding(
        workbook_path=path.resolve(),
        workbook_sha256=actual_sha,
        candidate_sheet=CANDIDATE_SHEET,
        candidate_case_marker_cell="BD2",
        candidate_timestamp_range="X{}:X{}".format(
            CANDIDATE_DATA_START_ROW,
            candidate_end,
        ),
        reference_sheet_prefix="Daten_",
        reference_case_marker_cell="BC2",
        reference_timestamp_range="X{}:X{}".format(
            REFERENCE_DATA_START_ROW,
            reference_end,
        ),
        hour_count=HOUR_COUNT,
        series=series,
    )


def load_test2a_diagnostic_workbook_binding(
    workbook_path: Union[str, Path],
    expected_sha256: str = EXPECTED_TEST2_WORKBOOK_SHA256,
) -> Test2ADiagnosticWorkbookBinding:
    """Return the checksum- and header-qualified official 2E1 schema."""

    path = Path(workbook_path).resolve()
    return _load_cached(str(path), str(expected_sha256).lower())
