"""Technical, non-normative evaluation of SIA 4010 Test 2 diagnostic 2E1.

The official Test 2 workbook plots annual values and hourly distributions for
diagnostic case 2E1, but it does not define an acceptance band or a PASS/FAIL
rule for that diagnostic.  This module therefore:

* reads the reference-program values directly from the checksum-qualified
  workbook;
* derives transparent min/max technical envelopes and the existing histogram
  scatter band;
* compares candidate hourly results to those references; and
* always reports that no SIA acceptance criterion is available.

The technical alignment can support an engineering review of a fixed-closed VE
awning representation.  It can never, by itself, qualify the optical mapping,
the dynamic Test 2A controller or SIA compliance.
"""

from __future__ import annotations

import hashlib
import json
import math
from dataclasses import asdict, dataclass
from functools import lru_cache
from pathlib import Path
from typing import Any, Dict, Mapping, Optional, Sequence, Tuple, Union

from openpyxl import load_workbook

from ..exceptions import ConfigurationError
from .distribution_reference import (
    DistributionQuantity,
    build_distribution_bands,
)
from .frequency_distribution import (
    DistributionBand,
    DistributionStatus,
    compare_distribution,
    histogram_counts,
)
from .test2a_diagnostic_workbook import (
    EXPECTED_TEST2_WORKBOOK_SHA256,
    HOUR_COUNT,
    load_test2a_diagnostic_workbook_binding,
)


EVALUATION_SCHEMA_VERSION = "1.0"
REFERENCE_ONLY_STATUS = (
    "REFERENCE_DIAGNOSTIC_RECORDED_NO_ACCEPTANCE_CRITERION"
)
NOT_CHECKABLE_STATUS = "NOT_CHECKABLE"

TOTAL_GAIN_SERIES_ID = "hourly_room_solar_heat_gain_total"
TRANSMITTED_SERIES_ID = (
    "hourly_transmitted_solar_radiation_excluding_secondary"
)

_ANNUAL_REFERENCE_COLUMNS = {
    TOTAL_GAIN_SERIES_ID: ("E", "F", "G", "H", "I", "J", "K", "L"),
    TRANSMITTED_SERIES_ID: ("R", "S", "T", "U", "V", "W", "X", "Y"),
}
_ANNUAL_REFERENCE_ROW = 19
_PROGRAM_HEADER_ROW = 9
_PROGRAM_VARIANT_ROW = 10


def _sha256(path: Path) -> str:
    """Return the lowercase SHA-256 of one artifact."""

    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_json(path: Path, payload: Mapping[str, Any]) -> None:
    """Atomically write deterministic UTF-8 JSON."""

    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(
        json.dumps(
            payload,
            ensure_ascii=False,
            indent=2,
            sort_keys=True,
        )
        + "\n",
        encoding="utf-8",
    )
    temporary.replace(path)


def _number(value: Any) -> Optional[float]:
    """Return a finite float, excluding bools and workbook blanks."""

    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    result = float(value)
    return result if math.isfinite(result) else None


def _program_name(worksheet: Any, column: str) -> str:
    """Build the same auditable program label used by the summary chart."""

    family = str(
        worksheet["{}{}".format(column, _PROGRAM_HEADER_ROW)].value or ""
    ).strip()
    variant = str(
        worksheet["{}{}".format(column, _PROGRAM_VARIANT_ROW)].value or ""
    ).strip()
    return " ".join(item for item in (family, variant) if item)


@dataclass(frozen=True)
class DiagnosticAnnualReference:
    """Annual values plotted by the official workbook for one 2E1 quantity."""

    series_id: str
    unit: str
    program_values: Tuple[Tuple[str, float], ...]
    minimum: float
    maximum: float
    mean: float
    source_locator: str
    acceptance_criterion: Optional[str] = None

    def to_dict(self) -> Dict[str, Any]:
        """Return a JSON-safe technical reference."""

        payload = asdict(self)
        payload["program_values"] = [
            {"program": name, "value": value}
            for name, value in self.program_values
        ]
        return payload


@dataclass(frozen=True)
class Test2A2E1ReferenceDataset:
    """Official source data available for a non-scored 2E1 comparison."""

    workbook_path: Path
    workbook_sha256: str
    annual_references: Tuple[DiagnosticAnnualReference, ...]
    total_gain_distribution: DistributionBand
    acceptance_criterion_available: bool = False

    def to_dict(self) -> Dict[str, Any]:
        """Return the complete source-traced reference dataset."""

        return {
            "workbook_path": str(self.workbook_path),
            "workbook_sha256": self.workbook_sha256,
            "diagnostic_case_id": "2E1",
            "annual_references": [
                item.to_dict() for item in self.annual_references
            ],
            "total_gain_distribution": asdict(
                self.total_gain_distribution
            ),
            "acceptance_criterion_available": (
                self.acceptance_criterion_available
            ),
            "claim_guardrail": (
                "The min/max and histogram scatter are technical envelopes "
                "derived from plotted reference programs. The workbook does "
                "not define them as SIA 4010 acceptance criteria for 2E1."
            ),
        }

    def annual(self, series_id: str) -> DiagnosticAnnualReference:
        """Return one annual reference or fail closed."""

        for item in self.annual_references:
            if item.series_id == series_id:
                return item
        raise ConfigurationError(
            "2E1 annual reference is unavailable for {!r}".format(series_id)
        )


@dataclass(frozen=True)
class Test2A2E1DiagnosticEvaluation:
    """Fail-closed candidate comparison with no compliance promotion."""

    status: str
    technical_alignment: str
    candidate_series_ids: Tuple[str, ...]
    candidate_hour_count: int
    candidate_annual_kwh: Tuple[Tuple[str, float], ...]
    annual_alignment: Tuple[Tuple[str, str], ...]
    total_gain_distribution_status: str
    total_gain_out_of_band_bins: Tuple[int, ...]
    required_eight_series_complete: bool
    engineering_review_ready: bool
    acceptance_criterion_available: bool
    compliance_pass: bool
    optical_mapping_qualified: bool
    dynamic_control_qualified: bool
    source_workbook_sha256: str

    def to_dict(self) -> Dict[str, Any]:
        """Return an audit report that cannot be mistaken for validation."""

        return {
            "schema_version": EVALUATION_SCHEMA_VERSION,
            "diagnostic_case_id": "2E1",
            "status": self.status,
            "technical_alignment": self.technical_alignment,
            "candidate_series_ids": list(self.candidate_series_ids),
            "candidate_hour_count": self.candidate_hour_count,
            "candidate_annual_kwh": dict(self.candidate_annual_kwh),
            "annual_alignment": dict(self.annual_alignment),
            "total_gain_distribution_status": (
                self.total_gain_distribution_status
            ),
            "total_gain_out_of_band_bins": list(
                self.total_gain_out_of_band_bins
            ),
            "required_eight_series_complete": (
                self.required_eight_series_complete
            ),
            "engineering_review_ready": self.engineering_review_ready,
            "acceptance_criterion_available": (
                self.acceptance_criterion_available
            ),
            "compliance_pass": self.compliance_pass,
            "optical_mapping_qualified": self.optical_mapping_qualified,
            "dynamic_control_qualified": self.dynamic_control_qualified,
            "source_workbook_sha256": self.source_workbook_sha256,
            "claim_guardrail": (
                "A technical alignment is not a SIA PASS. The supplied "
                "workbook defines no 2E1 acceptance criterion. Optical mapping "
                "requires a separate reviewed qualification decision, and "
                "dynamic Test 2A control remains a separate gate."
            ),
        }


def _annual_reference(
    worksheet: Any,
    series_id: str,
) -> DiagnosticAnnualReference:
    """Read one set of plotted annual reference-program values."""

    values = []
    for column in _ANNUAL_REFERENCE_COLUMNS[series_id]:
        value = _number(
            worksheet[
                "{}{}".format(column, _ANNUAL_REFERENCE_ROW)
            ].value
        )
        # In this workbook, a zero in the diagnostic annual-summary chart is a
        # missing result from that reference program, not a physical 2E1 annual
        # value. The positive contributors are the charted reference dataset.
        if value is None or value <= 0.0:
            continue
        values.append((_program_name(worksheet, column), value))
    if len(values) < 2:
        raise ConfigurationError(
            "Official 2E1 annual reference {!r} has fewer than two "
            "contributing programs".format(series_id)
        )
    numeric = [value for _name, value in values]
    return DiagnosticAnnualReference(
        series_id=series_id,
        unit="kWh",
        program_values=tuple(values),
        minimum=min(numeric),
        maximum=max(numeric),
        mean=sum(numeric) / len(numeric),
        source_locator=(
            "Resultaterfassung_Test2.xlsx!Zusammenfassung!{}19:{}19"
            .format(
                _ANNUAL_REFERENCE_COLUMNS[series_id][0],
                _ANNUAL_REFERENCE_COLUMNS[series_id][-1],
            )
        ),
        acceptance_criterion=None,
    )


@lru_cache(maxsize=8)
def _load_reference_cached(
    workbook_path: str,
    expected_sha256: str,
) -> Test2A2E1ReferenceDataset:
    """Load the official diagnostic references once per workbook identity."""

    path = Path(workbook_path)
    binding = load_test2a_diagnostic_workbook_binding(
        path,
        expected_sha256=expected_sha256,
    )
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if "Zusammenfassung" not in workbook.sheetnames:
            raise ConfigurationError(
                "Official Test 2 workbook is missing Zusammenfassung"
            )
        summary = workbook["Zusammenfassung"]
        annual = tuple(
            _annual_reference(summary, series_id)
            for series_id in (
                TOTAL_GAIN_SERIES_ID,
                TRANSMITTED_SERIES_ID,
            )
        )
    finally:
        workbook.close()

    label = "Solarer Wärmeeintrag gesamt"
    bands = build_distribution_bands(
        path,
        ("2 E1",),
        (
            DistributionQuantity(
                header_label=label,
                legend_key=label,
            ),
        ),
    )
    band = bands.get(("2 E1", label))
    if band is None or band.program_count < 2:
        raise ConfigurationError(
            "Official 2E1 total-gain reference distribution is unavailable"
        )
    return Test2A2E1ReferenceDataset(
        workbook_path=path.resolve(),
        workbook_sha256=binding.workbook_sha256,
        annual_references=annual,
        total_gain_distribution=band,
    )


def load_test2a_2e1_reference_dataset(
    workbook_path: Union[str, Path],
    expected_sha256: str = EXPECTED_TEST2_WORKBOOK_SHA256,
) -> Test2A2E1ReferenceDataset:
    """Return the checksum-qualified official 2E1 technical references."""

    return _load_reference_cached(
        str(Path(workbook_path).resolve()),
        str(expected_sha256).lower(),
    )


def _validated_series(
    series_id: str,
    values: Sequence[Any],
) -> Tuple[float, ...]:
    """Require exactly 8,760 finite numeric hourly values."""

    if isinstance(values, (str, bytes)) or len(values) != HOUR_COUNT:
        raise ConfigurationError(
            "2E1 series {!r} must contain exactly {} hourly values".format(
                series_id, HOUR_COUNT
            )
        )
    result = []
    for index, value in enumerate(values):
        number = _number(value)
        if number is None:
            raise ConfigurationError(
                "2E1 series {!r} contains a non-finite/non-numeric value at "
                "hour {}".format(series_id, index + 1)
            )
        result.append(number)
    return tuple(result)


def evaluate_test2a_2e1_diagnostic(
    candidate_series: Mapping[str, Sequence[Any]],
    reference: Test2A2E1ReferenceDataset,
    required_series_ids: Sequence[str],
) -> Test2A2E1DiagnosticEvaluation:
    """Compare candidate hourly series without inventing SIA acceptance."""

    normalized = {
        str(series_id): _validated_series(str(series_id), values)
        for series_id, values in candidate_series.items()
    }
    available = tuple(sorted(normalized))
    output_ids = tuple(
        series_id
        for series_id in (TOTAL_GAIN_SERIES_ID, TRANSMITTED_SERIES_ID)
        if series_id in normalized
    )
    if not output_ids:
        return Test2A2E1DiagnosticEvaluation(
            status=NOT_CHECKABLE_STATUS,
            technical_alignment="NOT_CHECKABLE",
            candidate_series_ids=available,
            candidate_hour_count=HOUR_COUNT if normalized else 0,
            candidate_annual_kwh=(),
            annual_alignment=(),
            total_gain_distribution_status=(
                DistributionStatus.NOT_CHECKABLE
            ),
            total_gain_out_of_band_bins=(),
            required_eight_series_complete=False,
            engineering_review_ready=False,
            acceptance_criterion_available=False,
            compliance_pass=False,
            optical_mapping_qualified=False,
            dynamic_control_qualified=False,
            source_workbook_sha256=reference.workbook_sha256,
        )

    annual_values = []
    annual_alignment = []
    for series_id in output_ids:
        value = sum(normalized[series_id]) / 1000.0
        band = reference.annual(series_id)
        within = band.minimum <= value <= band.maximum
        annual_values.append((series_id, value))
        annual_alignment.append(
            (
                series_id,
                "WITHIN_REFERENCE_ENVELOPE"
                if within
                else "OUTSIDE_REFERENCE_ENVELOPE",
            )
        )

    if TOTAL_GAIN_SERIES_ID in normalized:
        counts = histogram_counts(
            normalized[TOTAL_GAIN_SERIES_ID],
            reference.total_gain_distribution.upper_edges,
            include_overflow=reference.total_gain_distribution.include_overflow,
        )
        distribution = compare_distribution(
            counts,
            reference.total_gain_distribution,
        )
    else:
        distribution = compare_distribution(
            None,
            reference.total_gain_distribution,
        )

    annual_all_within = all(
        state == "WITHIN_REFERENCE_ENVELOPE"
        for _series_id, state in annual_alignment
    )
    if (
        TOTAL_GAIN_SERIES_ID in normalized
        and annual_all_within
        and DistributionStatus.is_passing(distribution.status)
    ):
        technical_alignment = "WITHIN_TECHNICAL_REFERENCE_ENVELOPE"
    elif TOTAL_GAIN_SERIES_ID not in normalized and annual_all_within:
        technical_alignment = "PARTIAL_ANNUAL_REFERENCE_ALIGNMENT"
    else:
        technical_alignment = "OUTSIDE_TECHNICAL_REFERENCE_ENVELOPE"

    required_complete = set(required_series_ids).issubset(normalized)
    engineering_review_ready = (
        required_complete
        and technical_alignment == "WITHIN_TECHNICAL_REFERENCE_ENVELOPE"
    )
    return Test2A2E1DiagnosticEvaluation(
        status=REFERENCE_ONLY_STATUS,
        technical_alignment=technical_alignment,
        candidate_series_ids=available,
        candidate_hour_count=HOUR_COUNT,
        candidate_annual_kwh=tuple(annual_values),
        annual_alignment=tuple(annual_alignment),
        total_gain_distribution_status=distribution.status,
        total_gain_out_of_band_bins=distribution.out_of_band_bins,
        required_eight_series_complete=required_complete,
        engineering_review_ready=engineering_review_ready,
        acceptance_criterion_available=False,
        compliance_pass=False,
        optical_mapping_qualified=False,
        dynamic_control_qualified=False,
        source_workbook_sha256=reference.workbook_sha256,
    )


def write_test2a_2e1_diagnostic_evaluation(
    output_path: Union[str, Path],
    evaluation: Test2A2E1DiagnosticEvaluation,
    reference: Test2A2E1ReferenceDataset,
) -> Path:
    """Write the technical evaluation and checksum as auditable artifacts."""

    path = Path(output_path)
    payload = evaluation.to_dict()
    payload["reference_dataset"] = reference.to_dict()
    _write_json(path, payload)
    checksum = _sha256(path)
    path.with_suffix(path.suffix + ".sha256").write_text(
        "{}  {}\n".format(checksum, path.name),
        encoding="ascii",
    )
    return path

