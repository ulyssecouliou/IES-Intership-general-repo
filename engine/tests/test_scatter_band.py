# -*- coding: utf-8 -*-
"""The implemented scatter band must reproduce the one the SIA computes itself.

Strategy: wherever an official workbook TABULATES a band (a "Mittelwert" /
"obere Grenze" / "untere Grenze" triplet), recompute that band with
`engine.scatter_band` from only the cells the workbook's own AVERAGE formula
designates, then require an exact match.

This validates three things at once that cannot be validated separately:
  1. the `mean +/- max|deviation|` formula;
  2. the extraction of the contributing set from the formula;
  3. the zero-floor handling (present in Test 1, absent in Test 2).

It also locks down the arbitration in
`traceability/streubereich-distributions.spec.md`: the band retained for the
frequency distributions is exactly the one the SIA applies where it does tabulate
it. If this test passes, the formula is no longer an interpretation, it is a
reproduction.

Checks that need the workbooks are skipped when the frozen source is absent --
see the CI blind spot documented in ADR-001 section 7bis.
"""

import os

import pytest

try:
    import openpyxl
except ImportError:
    openpyxl = None

from engine.scatter_band import (
    VERDICT_FAIL,
    VERDICT_NOT_CHECKABLE,
    VERDICT_PASS,
    VERDICT_PASS_WITH_RESERVATION,
    build_band,
    build_envelope,
    contributors_from_formula,
    is_passing,
    verdict,
)

_HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(_HERE, os.pardir, os.pardir))
WORKBOOK_DIR = os.environ.get(
    'SIA_WORKBOOK_DIR', os.path.join(ROOT, 'SIA_4010_geteilter_Link'))

# (test id, relative path, summary sheet). Sheet names stay verbatim: they are
# the German names inside the official files, i.e. data, not prose.
WORKBOOKS = [
    ('1', os.path.join('Test1', 'Resultaterfassung_Test1.xlsx'),
     u'Zusammenfassung Testf\xe4lle'),
    ('2', os.path.join('Test2', 'Resultaterfassung_Test2.xlsx'), 'Zusammenfassung'),
    ('5', os.path.join('Test5', 'Resultaterfassung_Test5.xlsx'), 'Zusammenfassung'),
]

TOLERANCE = 1e-9
MAX_ROWS = 400
MAX_COLS = 250


# --------------------------------------------------------------------------
# Pure checks -- no workbook needed, always run
# --------------------------------------------------------------------------

def test_basic_formula():
    """Explicit arithmetic case, computed by hand."""
    band = build_band([10.0, 12.0, 14.0, 20.0])
    assert abs(band.mean - 14.0) < TOLERANCE
    assert abs(band.max_deviation - 6.0) < TOLERANCE     # |20 - 14|
    assert abs(band.upper_bound - 20.0) < TOLERANCE
    assert abs(band.lower_bound - 8.0) < TOLERANCE       # 14 - 6, below min (10)


def test_zero_floor_is_not_applied_by_default():
    """The floor is a property of the table, never of the formula."""
    without = build_band([1.0, 2.0, 9.0])
    assert without.lower_bound < 0.0
    with_floor = build_band([1.0, 2.0, 9.0], floor_at_zero=True)
    assert with_floor.lower_bound == 0.0
    assert abs(with_floor.upper_bound - without.upper_bound) < TOLERANCE


def test_envelope_always_inside_the_band():
    """The arithmetic property the arbitration rests on: the envelope is stricter.

    max|v - mean| >= max - mean and >= mean - min, so [min, max] is always
    contained in [mean -/+ max_deviation].
    """
    datasets = [
        [1.0, 2.0, 3.0, 4.0],
        [0.0, 0.0, 0.0],
        [-5.0, 3.0, 100.0],
        [7.5],
        [1e-9, 2e-9, 3e-9],
        [1000.0, 1000.1],
    ]
    for values in datasets:
        band = build_band(values)
        low, high = build_envelope(values)
        assert band.lower_bound <= low + TOLERANCE, values
        assert band.upper_bound >= high - TOLERANCE, values


def test_absent_program_is_excluded_never_counted_as_zero():
    """A program that did not deliver the case drops out of the mean."""
    with_none = build_band([10.0, None, 20.0])
    without = build_band([10.0, 20.0])
    assert with_none.contributors == 2
    assert abs(with_none.mean - without.mean) < TOLERANCE
    # Counting it as zero would give a mean of 10; check that we do not.
    assert abs(with_none.mean - 15.0) < TOLERANCE


def test_three_state_verdict():
    references = [10.0, 12.0, 14.0, 20.0]   # envelope [10, 20]; band [8, 20]
    assert verdict(15.0, references) == VERDICT_PASS
    assert verdict(9.0, references) == VERDICT_PASS_WITH_RESERVATION
    assert verdict(7.0, references) == VERDICT_FAIL
    assert verdict(None, references) == VERDICT_NOT_CHECKABLE
    assert verdict(15.0, []) == VERDICT_NOT_CHECKABLE


def test_a_reserved_pass_counts_as_passing():
    """Callers must use is_passing(), or a pass silently becomes a failure."""
    assert is_passing(VERDICT_PASS)
    assert is_passing(VERDICT_PASS_WITH_RESERVATION)
    assert not is_passing(VERDICT_FAIL)
    assert not is_passing(VERDICT_NOT_CHECKABLE)


def test_bounds_are_inclusive():
    """The workbook defines no strict exclusion at the bounds."""
    band = build_band([10.0, 20.0])
    assert band.contains(band.lower_bound)
    assert band.contains(band.upper_bound)


def test_reading_the_contributing_set():
    """The two formula shapes found in the official workbooks."""
    assert contributors_from_formula('=AVERAGE(C16:F16)') == [
        'C16', 'D16', 'E16', 'F16']
    assert contributors_from_formula('=AVERAGE(E14,H14,K14,L14)') == [
        'E14', 'H14', 'K14', 'L14']
    assert contributors_from_formula('=AVERAGE(J9,N9,R9,V9)') == [
        'J9', 'N9', 'R9', 'V9']
    assert contributors_from_formula('=AVERAGE($C$16:$F$16)') == [
        'C16', 'D16', 'E16', 'F16']


def test_reading_refuses_rather_than_guesses():
    """Any unrecognised shape returns an empty list, never an approximation."""
    assert contributors_from_formula('=SUM(C16:F16)') == []
    assert contributors_from_formula('=AVERAGE(Daten_IDA!C16)') == []   # cross-sheet
    assert contributors_from_formula('=AVERAGE(IF(C16>0,C16))') == []   # nested
    assert contributors_from_formula('=AVERAGE(C16:F18)') == []         # multi-row
    assert contributors_from_formula(None) == []
    assert contributors_from_formula(42) == []


# --------------------------------------------------------------------------
# Confrontation with the official workbooks
# --------------------------------------------------------------------------

def _load(path, sheet_name):
    """Return (formulas, values): two dicts {coordinate: content}."""
    formulas = {}
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    try:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(min_row=1, max_row=MAX_ROWS,
                                       max_col=MAX_COLS):
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith('='):
                    formulas[cell.coordinate] = value
    finally:
        workbook.close()

    values = {}
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook[sheet_name]
        for row in worksheet.iter_rows(min_row=1, max_row=MAX_ROWS,
                                       max_col=MAX_COLS):
            for cell in row:
                if cell.value is not None:
                    values[cell.coordinate] = cell.value
    finally:
        workbook.close()
    return formulas, values


def _band_triplets(formulas):
    """Locate mean / upper / lower triplets by their formulas.

    No column is hard-coded: start from the `=AVERAGE(...)` cells and look, on
    the same row, for the two cells referencing them inside a `+MAX(ABS(` and a
    `-MAX(ABS(`.
    """
    import re
    means = dict((coord, f) for coord, f in formulas.items()
                 if f.upper().replace(' ', '').startswith('=AVERAGE('))
    triplets = []
    for mean_coord in means:
        row = re.match(r'^[A-Z]{1,3}(\d+)$', mean_coord).group(1)
        upper = lower = None
        for coord, formula in formulas.items():
            if re.match(r'^[A-Z]{1,3}' + row + r'$', coord) is None:
                continue
            compact = formula.upper().replace(' ', '')
            if mean_coord.upper() not in compact or 'MAX(ABS(' not in compact:
                continue
            if ('+MAX(ABS(' in compact) and upper is None:
                upper = coord
            elif ('-MAX(ABS(' in compact) and lower is None:
                lower = coord
        if upper is not None and lower is not None:
            floored = 'MAX(0,' in formulas[lower].upper().replace(' ', '')
            triplets.append((mean_coord, upper, lower, floored))
    return triplets


@pytest.mark.parametrize('test_id,relative,sheet_name', WORKBOOKS,
                         ids=[w[0] for w in WORKBOOKS])
def test_reproduces_the_tabulated_bands(test_id, relative, sheet_name):
    """The implemented formula exactly reproduces the SIA workbook's own bands."""
    if openpyxl is None:
        pytest.skip('openpyxl missing: install engine/requirements.txt')
    path = os.path.join(WORKBOOK_DIR, relative)
    if not os.path.exists(path):
        pytest.skip('Source workbook absent: ' + path +
                    ' -- reproduction of the official bands is NOT verified.')

    formulas, values = _load(path, sheet_name)
    triplets = _band_triplets(formulas)
    assert triplets, ('No band triplet located in Test ' + test_id +
                      ': the location strategy has stopped working.')

    mismatches = []
    reproduced = 0
    for mean_coord, upper_coord, lower_coord, floored in triplets:
        contributors = contributors_from_formula(formulas[mean_coord])
        if not contributors:
            mismatches.append(mean_coord + ': unusable AVERAGE formula: '
                              + formulas[mean_coord][:70])
            continue
        raw = [values.get(c) for c in contributors]
        if any(not isinstance(v, (int, float)) for v in raw if v is not None):
            continue  # an Excel error cell (#DIV/0!): out of scope here
        band = build_band(raw, floor_at_zero=floored)
        if band is None:
            continue
        expected = [('mean', mean_coord, band.mean),
                    ('upper bound', upper_coord, band.upper_bound),
                    ('lower bound', lower_coord, band.lower_bound)]
        for label, coord, computed in expected:
            tabulated = values.get(coord)
            if not isinstance(tabulated, (int, float)):
                continue
            if abs(float(tabulated) - computed) > 1e-6:
                mismatches.append(
                    'Test {0} {1} {2}: computed {3!r}, workbook {4!r} '
                    '(contributors {5}, floored={6})'.format(
                        test_id, coord, label, computed, tabulated,
                        contributors, floored))
        reproduced += 1

    assert not mismatches, ('{0} mismatch(es) over {1} band(s):\n{2}'.format(
        len(mismatches), reproduced, '\n'.join(mismatches[:20])))
    assert reproduced > 0, 'No band was actually recomputed.'
    print('\nTest {0}: {1} tabulated band(s) reproduced exactly.'.format(
        test_id, reproduced))
