"""Lecture des sorties horaires IESVE du dossier ZOER (fichiers ``.xlsx``).

Ce module ne contient AUCUNE logique normative : il se borne à extraire des
séries horodatées propres depuis les classeurs `Zone X_C{1,2}_Ergebnisse...xlsx`,
selon le mapping relevé et validé sur les fichiers réels (cf. README §Schéma).

Règles de lecture :
- ``openpyxl`` en ``read_only=True, data_only=True``. Seul l'onglet ``Ergebnisse``
  est une source ; ``Zusammenfassung`` et ``Grafik`` sont dérivés par formules
  Excel non recalculées (``data_only`` y renvoie ``None``) — ne pas s'en servir.
- Appariement des séries **par horodatage** (libellé de jour propagé + heure),
  jamais par index de ligne : la série opérative est à HH:30, la série
  extérieure à HH:00, et un pas extérieur manque au tout début (voir README).
"""

from __future__ import annotations

import datetime as _dt
import os
from dataclasses import dataclass, field
from typing import Optional

import openpyxl

# --- Emplacement par défaut des fichiers fournis (racine du dépôt) -----------
# Les classeurs ne sont pas dupliqués dans la fixture (binaires volumineux).
_REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
DEFAULT_INPUT_DIR = os.path.join(_REPO_ROOT, "VE - Validation for Swiss Building Regs")

ZONE_AREAS_M2 = {1: 91, 2: 147, 3: 77}  # Luftvolumenstrom.xlsx onglet C1 (métadonnée)


def _c1_filename(zone: int) -> str:
    return f"Zone {zone}_C1_Ergebnisse_2035_RCP85_DRY.xlsx"


def _c2_filename(zone: int) -> str:
    return f"Zone {zone}_C2_Ergebnisse_2035_RCP85_DRY.xlsx"


@dataclass(frozen=True)
class HourStep:
    """Un pas horaire apparié.

    ``operative`` : Dry resultant temperature (°C), température opérative /
    empfundene Temperatur (sortie ``.aps``).
    ``exterior``  : Dry-bulb temperature (°C) au même jour/heure, ou ``None`` si
    le pas extérieur manque (premier pas de la série C1).
    ``people_gain`` : People gain (kW), présent seulement pour le cas C2.
    """

    day_label: str          # p.ex. "Wed, 16/Apr" (sans année — étiquette DRY)
    hour: int               # heure de la journée 0..23 (issue du HH:30 opératif)
    operative: float
    exterior: Optional[float]
    people_gain: Optional[float] = None

    @property
    def occupied(self) -> bool:
        # Occupation C2 : People gain > 0 (cf. README §Occupation — seuil strict,
        # justifié par l'absence de valeurs dans (0 ; 0.05] kW).
        return self.people_gain is not None and self.people_gain > 0.0


@dataclass(frozen=True)
class ZoneCase:
    zone: int
    case: str                       # "C1" ou "C2"
    steps: list = field(default_factory=list)  # list[HourStep]

    @property
    def operative(self):
        return [s.operative for s in self.steps]

    @property
    def n_occupied(self) -> int:
        return sum(1 for s in self.steps if s.occupied)


def _hour_of(time_value) -> Optional[int]:
    """Extrait l'heure entière d'une cellule de temps openpyxl.

    Les cellules peuvent être ``datetime.time`` (HH:30 ou HH:00),
    ``datetime.datetime`` (00:00) ou ``datetime.timedelta`` (pas dégénéré du
    tout premier extérieur) — ce dernier est rejeté (``None``).
    """
    if isinstance(time_value, _dt.time):
        return time_value.hour
    if isinstance(time_value, _dt.datetime):
        return time_value.hour
    return None


def _propagate(day, current):
    """Propage vers le bas le libellé de jour épars des colonnes Date."""
    if isinstance(day, str) and day.strip():
        return day.strip()
    return current


def _date_key(label: Optional[str]) -> Optional[str]:
    """Clé de date robuste : la portion « JJ/Mois », sans le jour de semaine.

    Les deux blocs d'un même classeur n'ont PAS le même calendrier de jours de
    semaine (bloc extérieur « Thu, 16/Apr » vs bloc opératif « Wed, 16/Apr » —
    décalage d'un jour de semaine), mais la DATE (« 16/Apr ») concorde sur les
    183 jours. On apparie donc sur la date, jamais sur le libellé complet ni
    sur l'index de ligne.
    """
    if not label:
        return None
    return label.split(",")[-1].strip()


def _load_ergebnisse(path: str):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        assert "Ergebnisse" in wb.sheetnames, f"onglet Ergebnisse absent de {path}"
        ws = wb["Ergebnisse"]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()
    return rows


def load_c1(zone: int, input_dir: str = DEFAULT_INPUT_DIR) -> ZoneCase:
    """Charge un cas C1 (bauliche Grundanforderungen).

    Mapping ``Ergebnisse`` (0-based), en-têtes lignes 1-3, données dès la ligne 4 :
      col D idx 3  -> opérative (Dry resultant temperature, .aps)
      col A idx 0  -> Date opérative (épars)      / col B idx 1 -> Time HH:30
      col O idx 14 -> extérieure (Dry-bulb, .epw)
      col M idx 12 -> Date extérieure (épars)     / col N idx 13 -> Time HH:00
    """
    path = os.path.join(input_dir, _c1_filename(zone))
    rows = _load_ergebnisse(path)

    # Validation d'en-tête (§Schéma) — surface immédiate si le format change.
    h1 = rows[0]
    assert h1[3] and "Dry resultant temperature" in str(h1[3]), "C1: col D attendue = Dry resultant temperature"
    assert h1[14] and "Dry-bulb temperature" in str(h1[14]), "C1: col O attendue = Dry-bulb temperature"

    # Série extérieure indexée par (jour, heure) — appariement par horodatage.
    ext_by_key: dict = {}
    ext_day = None
    for r in rows[3:]:
        ext_day = _propagate(r[12] if len(r) > 12 else None, ext_day)
        hour = _hour_of(r[13] if len(r) > 13 else None)
        val = r[14] if len(r) > 14 else None
        if hour is not None and isinstance(val, (int, float)):
            ext_by_key[(_date_key(ext_day), hour)] = float(val)

    steps = []
    op_day = None
    for r in rows[3:]:
        op_day = _propagate(r[0], op_day)
        hour = _hour_of(r[1])
        op = r[3]
        if hour is None or not isinstance(op, (int, float)):
            continue
        exterior = ext_by_key.get((_date_key(op_day), hour))
        steps.append(HourStep(day_label=op_day, hour=hour, operative=float(op), exterior=exterior))

    return ZoneCase(zone=zone, case="C1", steps=steps)


def load_c2(zone: int, input_dir: str = DEFAULT_INPUT_DIR) -> ZoneCase:
    """Charge un cas C2 (thermischer Komfort).

    Mapping ``Ergebnisse`` (0-based), marqueurs "Paste below" ligne 1, en-têtes
    lignes 2-4, données dès la ligne 5 :
      col C idx 2 -> extérieure (Dry-bulb, .epw)   / col A,B idx 0,1 -> Date/Time HH:00
      col F idx 5 -> opérative (Dry resultant, .aps)/ col D,E idx 3,4 -> Date/Time HH:30
      col I idx 8 -> People gain (kW, .aps)          / col G,H idx 6,7 -> Date/Time HH:30
    """
    path = os.path.join(input_dir, _c2_filename(zone))
    rows = _load_ergebnisse(path)

    hdr = rows[3]  # ligne Excel 4 : "Date/Time/source"
    assert any(cell and "SMA" in str(cell) for cell in hdr), "C2: source .epw attendue en ligne d'en-tête"

    # Extérieur par (jour, heure).
    ext_by_key: dict = {}
    ext_day = None
    for r in rows[4:]:
        ext_day = _propagate(r[0], ext_day)
        hour = _hour_of(r[1])
        val = r[2] if len(r) > 2 else None
        if hour is not None and isinstance(val, (int, float)):
            ext_by_key[(_date_key(ext_day), hour)] = float(val)

    steps = []
    op_day = None
    for r in rows[4:]:
        op_day = _propagate(r[3], op_day)
        hour = _hour_of(r[4])
        op = r[5]
        pg = r[8] if len(r) > 8 else None
        if hour is None or not isinstance(op, (int, float)):
            continue
        exterior = ext_by_key.get((_date_key(op_day), hour))
        people = float(pg) if isinstance(pg, (int, float)) else None
        steps.append(HourStep(day_label=op_day, hour=hour, operative=float(op),
                              exterior=exterior, people_gain=people))

    return ZoneCase(zone=zone, case="C2", steps=steps)
