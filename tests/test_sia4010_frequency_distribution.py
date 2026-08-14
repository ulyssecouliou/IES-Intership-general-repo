"""Tests for the SIA 4010 frequency-distribution criterion.

The pure-Python engine is proven on the REAL official legend and a real
reference hourly column (count conservation), plus synthetic scatter bands to
exercise PASS / FAIL / NOT_CHECKABLE. Nothing is fabricated: bins come verbatim
from the legend and the band comes only from the reference programs' counts.
"""

import glob
import os
import unittest
import warnings
from pathlib import Path

from openpyxl import Workbook

from swiss_sia.reference_model.sia4010.distribution_reference import (
    DISTRIBUTION_CRITERIA,
    DISTRIBUTION_WORKBOOK_EVIDENCE,
    DistributionQuantity,
    build_distribution_bands,
)
from swiss_sia.reference_model.sia4010.frequency_distribution import (
    DistributionStatus,
    build_scatter_band,
    compare_distribution,
    histogram_counts,
    parse_frequency_classes,
)

REPO_ROOT = Path(__file__).resolve().parents[1]
TEST2_WB = REPO_ROOT / "SIA_4010_geteilter_Link" / "Test2" / "Resultaterfassung_Test2.xlsx"
TEST3_WB = REPO_ROOT / "SIA_4010_geteilter_Link" / "Test3" / "Resultaterfassung_Test3.xlsx"
HAS_TEST2 = TEST2_WB.is_file()
HAS_TEST3 = TEST3_WB.is_file()
# The real Test 2 extraction reads several ~8760-row reference sheets (~1 min);
# keep it out of the fast suite unless explicitly requested.
RUN_HEAVY = bool(os.environ.get("SIA4010_RUN_HEAVY"))
SYNTHETIC_WB = REPO_ROOT / ".codex_tmp" / "sia4010_freq_dist" / "synthetic.xlsx"


def _build_synthetic_workbook(path: Path) -> Path:
    """Build a tiny workbook mirroring the Daten_* + legend layout for fast tests."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    legend = workbook.active
    legend.title = "Haeufigkeitsklassen"
    legend["A1"] = "Klassen"
    legend["B2"] = "Q1"
    legend["B3"] = "W"
    legend["A4"], legend["B4"] = 1, 10
    legend["A5"], legend["B5"] = 2, 100

    def add_daten(title, vals_1a, vals_1b):
        """Add a Daten_* sheet: header row with 1A/1B markers and Q1 columns."""
        worksheet = workbook.create_sheet(title)
        worksheet.cell(row=1, column=1, value="Programm")
        worksheet.cell(row=1, column=2, value="1A")
        worksheet.cell(row=1, column=3, value="Q1 (W)")
        worksheet.cell(row=1, column=4, value="Other (W)")
        worksheet.cell(row=1, column=5, value="1B")
        worksheet.cell(row=1, column=6, value="Q1 (W)")
        worksheet.cell(row=1, column=7, value="Other (W)")
        for offset in range(max(len(vals_1a), len(vals_1b))):
            if offset < len(vals_1a):
                worksheet.cell(row=2 + offset, column=3, value=vals_1a[offset])
            if offset < len(vals_1b):
                worksheet.cell(row=2 + offset, column=6, value=vals_1b[offset])

    add_daten("Daten_ProgA", [5, 50, 500], [50])
    add_daten("Daten_ProgB", [5, 5, 500], [])  # ProgB did not provide case 1B
    add_daten("Daten_Testprogramm", [999, 999, 999], [999])  # candidate, must be excluded
    workbook.save(path)
    return path


class HistogramEngineTests(unittest.TestCase):
    """Excel-FREQUENCY binning, band building and comparison, pure Python."""

    def test_histogram_uses_excel_frequency_semantics(self):
        # edges [10, 100, 150]: class0 v<=10, class1 10<v<=100, class2 100<v<=150,
        # overflow v>150. Values chosen to land one in each, incl. a boundary hit.
        counts = histogram_counts([5, 10, 15, 150, 1000, None, "x"], [10, 100, 150])
        self.assertEqual(counts, (2, 1, 1, 1))  # 5 and 10 -> class0; 15 -> c1; 150 -> c2; 1000 -> overflow

    def test_official_displayed_classes_keep_overflow_separate(self):
        counts = histogram_counts(
            [5, 10, 15, 150, 1000],
            [10, 100, 150],
            include_overflow=False,
        )
        self.assertEqual(counts, (2, 1, 1))
        self.assertEqual(sum(counts), 4)

    def test_build_scatter_band_is_the_per_bin_min_max_envelope(self):
        """The written authority clarification makes min/max the acceptance band."""

        band = build_scatter_band(
            [(10, 5, 2), (8, 7, 3), (12, 4, 1)],
            quantity="Q",
            unit="W",
            upper_edges=[10, 100],
            source_locator="loc",
        )
        self.assertIsNotNone(band)
        for index, expected in enumerate((8.0, 4.0, 1.0)):
            self.assertAlmostEqual(band.lower_counts[index], expected, places=9)
        for index, expected in enumerate((12.0, 7.0, 3.0)):
            self.assertAlmostEqual(band.upper_counts[index], expected, places=9)
        self.assertEqual(band.program_count, 3)

    def test_build_scatter_band_records_envelope_and_old_diagnostic(self):
        """The accepted envelope and superseded symmetric diagnostic are auditable."""

        band = build_scatter_band(
            [(10, 5, 2), (8, 7, 3), (12, 4, 1)],
            quantity="Q",
            unit="W",
            upper_edges=[10, 100],
            source_locator="loc",
        )
        self.assertEqual(band.envelope_lower_counts, (8.0, 4.0, 1.0))
        self.assertEqual(band.envelope_upper_counts, (12.0, 7.0, 3.0))
        for index, expected in enumerate((8.0, 11.0 / 3.0, 1.0)):
            self.assertAlmostEqual(
                band.symmetric_lower_counts[index], expected, places=9
            )
        self.assertEqual(band.symmetric_upper_counts, (12.0, 7.0, 3.0))

    def test_accepted_envelope_is_inside_the_old_diagnostic_band(self):
        """The superseded symmetric diagnostic remains wider or equal."""

        for programs in (
            [(1, 2, 3), (4, 5, 6), (7, 8, 9)],
            [(0, 0, 0), (10, 0, 0)],
            [(5, 5, 5), (5, 5, 5)],
            [(0, 100, 3), (2, 1, 0), (7, 7, 7)],
        ):
            band = build_scatter_band(
                programs, quantity="Q", unit="W", upper_edges=[10, 100],
                source_locator="loc",
            )
            for index in range(len(band.lower_counts)):
                self.assertLessEqual(
                    band.symmetric_lower_counts[index],
                    band.lower_counts[index] + 1e-9,
                    msg="{} bin {}".format(programs, index),
                )
                self.assertGreaterEqual(
                    band.symmetric_upper_counts[index],
                    band.upper_counts[index] - 1e-9,
                    msg="{} bin {}".format(programs, index),
                )

    def test_build_scatter_band_rejects_mismatched_bin_lengths(self):
        band = build_scatter_band(
            [(10, 5, 2), (1, 2)],  # second program has wrong length -> dropped
            quantity="Q", unit="W", upper_edges=[10, 100], source_locator="loc",
        )
        self.assertEqual(band.program_count, 1)

    def test_build_scatter_band_none_when_no_program_contributes(self):
        self.assertIsNone(
            build_scatter_band([], quantity="Q", unit="W", upper_edges=[10], source_locator="loc")
        )

    def _band(self):
        return build_scatter_band(
            [(10, 5, 2), (8, 7, 3), (12, 4, 1)],
            quantity="Q", unit="W", upper_edges=[10, 100], source_locator="loc",
        )

    def test_candidate_within_band_passes(self):
        outcome = compare_distribution((9, 6, 2), self._band())
        self.assertEqual(outcome.status, DistributionStatus.PASS)
        self.assertEqual(outcome.out_of_band_bins, ())

    def test_candidate_outside_band_fails_with_bin_indices(self):
        outcome = compare_distribution((13, 6, 0), self._band())  # bin0 high, bin2 low
        self.assertEqual(outcome.status, DistributionStatus.FAIL)
        self.assertEqual(outcome.out_of_band_bins, (0, 2))

    def test_missing_candidate_is_not_checkable(self):
        outcome = compare_distribution(None, self._band())
        self.assertEqual(outcome.status, DistributionStatus.NOT_CHECKABLE)

    def test_missing_band_is_not_checkable(self):
        outcome = compare_distribution((1, 2, 3), None)
        self.assertEqual(outcome.status, DistributionStatus.NOT_CHECKABLE)

    def test_bin_length_mismatch_is_not_checkable(self):
        outcome = compare_distribution((1, 2), self._band())  # band has 3 bins
        self.assertEqual(outcome.status, DistributionStatus.NOT_CHECKABLE)

    def test_inside_old_symmetric_band_but_outside_envelope_fails(self):
        """The authority-confirmed min/max envelope is now enforced."""

        outcome = compare_distribution((10, 3.8, 2), self._band())
        self.assertEqual(outcome.status, DistributionStatus.FAIL)
        self.assertEqual(outcome.out_of_band_bins, (1,))
        self.assertIn("min/max envelope", outcome.message)

    def test_historical_reserved_pass_still_deserializes_as_passing(self):
        """Backward compatibility does not change new comparison behavior."""

        self.assertTrue(DistributionStatus.is_passing(DistributionStatus.PASS))
        self.assertTrue(DistributionStatus.is_passing(
            DistributionStatus.PASS_WITH_RESERVATION))
        self.assertFalse(DistributionStatus.is_passing(DistributionStatus.FAIL))
        self.assertFalse(DistributionStatus.is_passing(
            DistributionStatus.NOT_CHECKABLE))

    def test_no_reservation_when_the_envelope_is_unknown(self):
        """A band built without envelope data never raises a false reservation."""

        from swiss_sia.reference_model.sia4010.frequency_distribution import (
            DistributionBand,
        )

        band = DistributionBand(
            quantity="Q", unit="W", upper_edges=(10.0, 100.0),
            lower_counts=(0.0, 0.0, 0.0), upper_counts=(20.0, 20.0, 20.0),
            program_count=3, source_locator="loc",
        )
        outcome = compare_distribution((1, 2, 3), band)
        self.assertEqual(outcome.status, DistributionStatus.PASS)


@unittest.skipUnless(HAS_TEST2, "official SIA 4010 Test 2 workbook not present")
class RealLegendAndDataTests(unittest.TestCase):
    """Prove the parser and histogram against the real Test 2 workbook."""

    def test_parses_official_class_legend(self):
        classes = parse_frequency_classes(TEST2_WB)
        self.assertIn("Solarer Wärmeeintrag gesamt", classes)
        solar = classes["Solarer Wärmeeintrag gesamt"]
        self.assertEqual(solar.unit, "W")
        self.assertEqual(solar.upper_edges[:3], (10.0, 100.0, 150.0))

    def test_histogram_conserves_count_on_a_real_reference_column(self):
        import openpyxl

        warnings.filterwarnings("ignore")
        classes = parse_frequency_classes(TEST2_WB)
        edges = classes["Solarer Wärmeeintrag gesamt"].upper_edges
        workbook = openpyxl.load_workbook(TEST2_WB, data_only=True, read_only=True)
        try:
            worksheet = workbook["Daten_Excel"]
            rows = list(worksheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        # Locate the header row and the first column naming the solar heat gain
        # (reference sheets carry headers on a different row than the candidate).
        header_row_idx = next(
            (
                index
                for index, row in enumerate(rows)
                if any(isinstance(v, str) and "Solarer Wärmeeintrag gesamt" in v for v in row)
            ),
            None,
        )
        self.assertIsNotNone(header_row_idx, "solar heat gain header row not found")
        header = rows[header_row_idx]
        col = next(
            index
            for index, value in enumerate(header)
            if isinstance(value, str) and "Solarer Wärmeeintrag gesamt" in value
        )
        series = [row[col] for row in rows[header_row_idx + 1:] if col < len(row)]
        numeric = [v for v in series if isinstance(v, (int, float)) and not isinstance(v, bool)]
        counts = histogram_counts(series, edges)
        self.assertEqual(len(counts), len(edges) + 1)
        self.assertEqual(sum(counts), len(numeric))  # every hour lands in exactly one class


class SyntheticExtractionTests(unittest.TestCase):
    """Cover build_distribution_bands quickly on a tiny synthetic workbook."""

    @classmethod
    def setUpClass(cls):
        """Build the synthetic workbook and extract its bands once."""
        _build_synthetic_workbook(SYNTHETIC_WB)
        cls.bands = build_distribution_bands(
            SYNTHETIC_WB,
            case_ids=["1A", "1B"],
            quantities=[DistributionQuantity(header_label="Q1", legend_key="Q1")],
        )

    def test_extracts_bands_for_both_cases(self):
        self.assertEqual(set(self.bands), {("1A", "Q1"), ("1B", "Q1")})

    def test_candidate_sheet_is_excluded_from_the_band(self):
        # ProgA and ProgB contribute 1A; the candidate Daten_Testprogramm must not.
        band = self.bands[("1A", "Q1")]
        self.assertEqual(band.program_count, 2)
        self.assertEqual(band.lower_counts, (1, 0))
        self.assertEqual(band.upper_counts, (2, 1))
        self.assertFalse(band.include_overflow)
        self.assertEqual(
            dict(band.outside_class_counts),
            {"Daten_ProgA": 1, "Daten_ProgB": 1},
        )
        self.assertEqual(band.unit, "W")

    def test_incomplete_reference_series_are_retained_and_flagged(self):
        """Authority instruction: do not repair or exclude supplied series."""

        band = self.bands[("1A", "Q1")]
        self.assertEqual(
            dict(band.reference_hour_totals),
            {"Daten_ProgA": 3, "Daten_ProgB": 3},
        )
        self.assertEqual(
            set(band.incomplete_reference_programs),
            {"Daten_ProgA", "Daten_ProgB"},
        )
        outcome = compare_distribution((2, 1), band)
        self.assertEqual(outcome.status, DistributionStatus.PASS)
        self.assertIn("retained unchanged", outcome.message)
        self.assertIn("outside the declared class boundaries", outcome.message)

    def test_not_provided_case_has_fewer_contributors(self):
        # Only ProgA provided case 1B (ProgB's 1B column is empty).
        band = self.bands[("1B", "Q1")]
        self.assertEqual(band.program_count, 1)
        self.assertEqual(band.lower_counts, (0, 1))

    def test_comparator_against_synthetic_band(self):
        band = self.bands[("1A", "Q1")]
        self.assertEqual(
            compare_distribution((2, 1), band).status, DistributionStatus.PASS
        )
        self.assertEqual(
            compare_distribution((3, 0), band).status, DistributionStatus.FAIL
        )

    def test_evaluate_distribution_criteria_fail_closed_then_pass(self):
        from swiss_sia.reference_model.sia4010.distribution_reference import (
            evaluate_distribution_criteria,
        )

        quantities = [DistributionQuantity(header_label="Q1", legend_key="Q1")]
        # No candidate supplied -> every (case, quantity) is NOT_CHECKABLE.
        without = evaluate_distribution_criteria(SYNTHETIC_WB, ["1A", "1B"], quantities)
        self.assertTrue(without)
        self.assertTrue(all(o.status == DistributionStatus.NOT_CHECKABLE for o in without))
        # Candidate inside the band -> PASS.
        band_1a = self.bands[("1A", "Q1")]
        candidate = {
            ("1A", "Q1"): tuple(
                (lo + hi) // 2 for lo, hi in zip(band_1a.lower_counts, band_1a.upper_counts)
            ),
            ("1B", "Q1"): self.bands[("1B", "Q1")].lower_counts,
        }
        withc = evaluate_distribution_criteria(
            SYNTHETIC_WB, ["1A", "1B"], quantities, candidate_distributions=candidate
        )
        self.assertEqual({o.status for o in withc}, {DistributionStatus.PASS})


@unittest.skipUnless(
    HAS_TEST2 and RUN_HEAVY,
    "set SIA4010_RUN_HEAVY=1 (with the official package) to run the heavy real extraction",
)
class Test2ReferenceBandTests(unittest.TestCase):
    """Build the real Test 2 solar-heat-gain scatter band from the reference sheets."""

    SOLAR = "Solarer Wärmeeintrag gesamt"

    @classmethod
    def setUpClass(cls):
        """Extract the reference bands once (reads several large data sheets)."""
        warnings.filterwarnings("ignore")
        cls.bands = build_distribution_bands(
            TEST2_WB,
            case_ids=["2A", "2B", "2C", "2D"],
            quantities=[DistributionQuantity(header_label=cls.SOLAR, legend_key=cls.SOLAR)],
        )

    def test_bands_built_for_all_scored_cases(self):
        self.assertEqual({case for case, _ in self.bands}, {"2A", "2B", "2C", "2D"})

    def test_contributing_program_counts_match_the_workbook(self):
        # Derived by inspecting which reference Daten_* sheets provided each case.
        counts = {
            case: self.bands[(case, self.SOLAR)].program_count
            for case in ["2A", "2B", "2C", "2D"]
        }
        self.assertEqual(counts, {"2A": 6, "2B": 4, "2C": 5, "2D": 4})

    def test_band_shape_is_consistent(self):
        band = self.bands[("2A", self.SOLAR)]
        self.assertEqual(band.unit, "W")
        self.assertFalse(band.include_overflow)
        self.assertEqual(len(band.lower_counts), len(band.upper_edges))
        self.assertEqual(
            set(dict(band.reference_hour_totals).values()),
            {8760},
        )
        self.assertTrue(
            all(count >= 0 for count in dict(band.outside_class_counts).values())
        )
        self.assertTrue(
            all(lo <= hi for lo, hi in zip(band.lower_counts, band.upper_counts))
        )

    def test_synthetic_midpoint_candidate_passes_real_band(self):
        # Candidate = per-bin midpoint of the real band: clearly synthetic (not a
        # real VE result), lies inside the band by construction -> PASS.
        band = self.bands[("2A", self.SOLAR)]
        candidate = tuple(
            (lo + hi) // 2 for lo, hi in zip(band.lower_counts, band.upper_counts)
        )
        outcome = compare_distribution(candidate, band)
        self.assertEqual(outcome.status, DistributionStatus.PASS)

    def test_candidate_outside_real_band_fails(self):
        band = self.bands[("2A", self.SOLAR)]
        candidate = tuple(hi + 100 for hi in band.upper_counts)
        outcome = compare_distribution(candidate, band)
        self.assertEqual(outcome.status, DistributionStatus.FAIL)
        self.assertTrue(outcome.out_of_band_bins)


class DistributionRegistryTests(unittest.TestCase):
    """Pin which tests carry a registered distribution criterion, and their layout."""

    def test_registered_tests_cover_every_spec_requirement(self):
        # Spezifikation_Test2/3/5 are the only specs that require the hourly
        # frequency distribution to lie within the reference scatter band.
        self.assertEqual(set(DISTRIBUTION_CRITERIA), {"2", "3", "5"})

    def test_tests_4_6_distribution_sheets_are_no_longer_denied(self):
        """Presence is proven; exact pass-gate scope remains fail-closed."""

        self.assertEqual(set(DISTRIBUTION_WORKBOOK_EVIDENCE), {"4", "6"})
        for test_id, evidence in DISTRIBUTION_WORKBOOK_EVIDENCE.items():
            with self.subTest(test_id=test_id):
                self.assertEqual(evidence["legend_sheet"], "Haeufigkeitskassen")
                self.assertEqual(
                    evidence["status"],
                    "REQUIRED_OUTPUT_SCOPE_CONFIRMED_ACCEPTANCE_GATE_PENDING",
                )
                self.assertGreaterEqual(len(evidence["legend_quantities"]), 10)

    def test_every_validation_class_has_its_criteria_registered(self):
        # A class is only validatable when every test family it requires has both
        # a band parser and, where the spec demands it, a distribution criterion.
        import re

        from swiss_sia.config import SIA4010_CLASS_TEST_MATRIX
        from swiss_sia.reference_model.sia4010.test_runner import _TEST_PARSERS

        distribution_required = {"2", "3", "5"}
        for class_id, variants in SIA4010_CLASS_TEST_MATRIX.items():
            families = {re.match(r"test_(\d)", v).group(1) for v in variants}
            with self.subTest(validation_class=class_id):
                self.assertTrue(families <= set(_TEST_PARSERS))
                self.assertEqual(
                    (families & distribution_required) - set(DISTRIBUTION_CRITERIA),
                    set(),
                )

    def test_test5_uses_the_split_header_layout(self):
        entry = DISTRIBUTION_CRITERIA["5"]
        self.assertEqual(entry["layout"], "split_header")
        self.assertEqual(entry["data_prefix"], "Daten ")  # space, not underscore
        self.assertEqual(entry["candidate_sheet"], "Daten Testprogramm")
        self.assertEqual(entry["scored_section"], "Testgrössen")
        self.assertEqual(entry["diagnostic_section"], "Diagnosegrössen")
        self.assertEqual(
            list(entry["case_ids"]), ["Test 5A", "Test 5B", "Test 5C", "Test 5D"]
        )
        labels = [q.header_label for q in entry["quantities"]]
        # Both source spellings of the WRG transfer quantities are registered;
        # the 5C sheets repeat a word in the official workbook.
        self.assertIn("WRG Übertragungsleistung gesamt", labels)
        self.assertIn("WRG Übertragungsleistungleistung gesamt", labels)
        # The humidifier power has no legend bin column, so it is not scored.
        self.assertNotIn("Leistung Befeuchter", labels)

    def test_test3_registry_entry_is_complete(self):
        entry = DISTRIBUTION_CRITERIA["3"]
        self.assertEqual(len(entry["case_ids"]), 12)
        self.assertEqual(entry["case_ids"][0], "3A")
        self.assertEqual(entry["case_ids"][-1], "3L")
        self.assertEqual(entry["data_prefix"], "Daten_")
        self.assertEqual(entry["candidate_sheet"], "Daten_Testprogramm")
        labels = [q.header_label for q in entry["quantities"]]
        self.assertEqual(labels, ["Beleuchtungsleistung"])
        # Illuminance is a Diagnosegroesse and must never be scored.
        self.assertNotIn("Beleuchtungsstärke", labels)


@unittest.skipUnless(
    HAS_TEST3 and RUN_HEAVY,
    "set SIA4010_RUN_HEAVY=1 (with the official package) to run the heavy real extraction",
)
class Test3ReferenceBandTests(unittest.TestCase):
    """Build the real Test 3 lighting-power scatter bands from the reference sheets."""

    @classmethod
    def setUpClass(cls):
        """Extract the Test 3 reference bands once."""
        warnings.filterwarnings("ignore")
        entry = DISTRIBUTION_CRITERIA["3"]
        cls.entry = entry
        cls.bands = build_distribution_bands(
            TEST3_WB,
            entry["case_ids"],
            entry["quantities"],
            candidate_sheet=entry["candidate_sheet"],
            data_prefix=entry["data_prefix"],
        )

    def test_every_scored_case_gets_a_band(self):
        self.assertEqual(
            {case for case, _ in self.bands}, set(self.entry["case_ids"])
        )

    def test_bands_are_in_watts_with_contributing_programs(self):
        for key, band in self.bands.items():
            with self.subTest(case=key[0]):
                self.assertEqual(band.unit, "W")
                self.assertGreaterEqual(band.program_count, 2)
                # Post-2026-08-10 convention (see Test2ReferenceBandTests): no
                # overflow bin, so counts align one-to-one with the edges;
                # out-of-class hours are tracked separately, not as a +1 bin.
                self.assertFalse(band.include_overflow)
                self.assertEqual(len(band.lower_counts), len(band.upper_edges))
                self.assertTrue(
                    all(lo <= hi for lo, hi in zip(band.lower_counts, band.upper_counts))
                )

    def test_midpoint_candidate_passes_and_outside_fails(self):
        band = self.bands[("3A", "Beleuchtungsleistung")]
        inside = tuple(
            (lo + hi) // 2 for lo, hi in zip(band.lower_counts, band.upper_counts)
        )
        self.assertEqual(
            compare_distribution(inside, band).status, DistributionStatus.PASS
        )
        outside = tuple(hi + 500 for hi in band.upper_counts)
        self.assertEqual(
            compare_distribution(outside, band).status, DistributionStatus.FAIL
        )


SPLIT_WB = REPO_ROOT / ".codex_tmp" / "sia4010_freq_dist" / "split_header.xlsx"


def _build_split_header_workbook(path: Path) -> Path:
    """Build a tiny workbook mirroring the Test 5 split-header layout.

    Row 1 carries the case markers and the section labels; row 3 carries the
    quantity labels; data starts at row 4. One reference sheet deliberately
    omits the diagnostic-section marker, reproducing the real workbook defect.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    legend = workbook.active
    legend.title = "Haeufigkeitsklassen"
    legend["A1"] = "Klassen"
    legend["B2"], legend["B3"] = "FanPower", "W"
    legend["A4"], legend["B4"] = 1, 10
    legend["A5"], legend["B5"] = 2, 100

    def add(title, scored_values, diag_values, with_boundary=True):
        """Add one data sheet: case marker + sections on row 1, labels on row 3."""
        sheet = workbook.create_sheet(title)
        sheet.cell(row=1, column=1, value="Resultate")
        sheet.cell(row=1, column=2, value="Case A")       # case marker
        sheet.cell(row=1, column=3, value="Testgroessen")  # scored section label
        if with_boundary:
            sheet.cell(row=1, column=4, value="Diag")     # diagnostic boundary
        sheet.cell(row=3, column=3, value="FanPower")     # scored quantity
        sheet.cell(row=3, column=4, value="Noise")        # diagnostic quantity
        for offset, value in enumerate(scored_values):
            sheet.cell(row=4 + offset, column=3, value=value)
        for offset, value in enumerate(diag_values):
            sheet.cell(row=4 + offset, column=4, value=value)

    add("Daten ProgA", [5, 50, 500], [1, 1, 1])
    add("Daten ProgB", [5, 5, 500], [1, 1, 1])
    # ProgC omits the boundary: its span would otherwise swallow "Noise".
    add("Daten ProgC", [50, 50, 50], [1, 1, 1], with_boundary=False)
    add("Daten Testprogramm", [999, 999, 999], [9, 9, 9])  # candidate, excluded
    workbook.save(path)
    return path


class SplitHeaderExtractionTests(unittest.TestCase):
    """Cover the Test 5 style extractor and its two safeguards."""

    @classmethod
    def setUpClass(cls):
        """Build the synthetic split-header workbook and extract its bands."""
        from swiss_sia.reference_model.sia4010.distribution_reference import (
            build_split_header_distribution_bands,
        )

        _build_split_header_workbook(SPLIT_WB)
        cls.bands = build_split_header_distribution_bands(
            SPLIT_WB,
            case_ids=["Case A"],
            quantities=[DistributionQuantity(header_label="FanPower", legend_key="FanPower")],
            scored_section="Testgroessen",
            diagnostic_section="Diag",
            candidate_sheet="Daten Testprogramm",
            data_prefix="Daten ",
        )

    def test_only_the_scored_quantity_is_banded(self):
        self.assertEqual(set(self.bands), {("Case A", "FanPower")})

    def test_diagnostic_column_never_enters_the_band(self):
        # "Noise" is a diagnostic quantity: it must not be registered at all,
        # even though ProgC's span lacks the diagnostic boundary.
        self.assertNotIn(("Case A", "Noise"), self.bands)

    def test_candidate_sheet_is_excluded_and_all_references_contribute(self):
        band = self.bands[("Case A", "FanPower")]
        self.assertEqual(band.program_count, 3)  # ProgA, ProgB, ProgC - not the candidate
        self.assertEqual(band.unit, "W")

    def test_band_is_the_per_bin_min_max_envelope_of_reference_programs(self):
        band = self.bands[("Case A", "FanPower")]
        # The official displayed classes are <=10 and <=100. The >100 Excel
        # overflow is audited separately, not scored as an invented class.
        for index, expected in enumerate((0.0, 0.0)):
            self.assertAlmostEqual(band.lower_counts[index], expected, places=9)
        for index, expected in enumerate((2.0, 3.0)):
            self.assertAlmostEqual(band.upper_counts[index], expected, places=9)
        # Backward-compatible envelope fields repeat the accepted interval.
        self.assertEqual(band.envelope_lower_counts, (0.0, 0.0))
        self.assertEqual(band.envelope_upper_counts, (2.0, 3.0))
        self.assertEqual(
            dict(band.outside_class_counts),
            {"Daten ProgA": 1, "Daten ProgB": 1, "Daten ProgC": 0},
        )
        for index, expected in enumerate((0.0, -1.0 / 3.0)):
            self.assertAlmostEqual(
                band.symmetric_lower_counts[index], expected, places=9
            )
        for index, expected in enumerate((2.0, 3.0)):
            self.assertAlmostEqual(
                band.symmetric_upper_counts[index], expected, places=9
            )

    def test_single_program_bands_are_refused(self):
        from swiss_sia.reference_model.sia4010.distribution_reference import (
            MIN_REFERENCE_PROGRAMS,
            build_split_header_distribution_bands,
        )

        self.assertGreaterEqual(MIN_REFERENCE_PROGRAMS, 2)
        # Only one reference sheet present -> no scatter band can be formed.
        lonely = REPO_ROOT / ".codex_tmp" / "sia4010_freq_dist" / "split_lonely.xlsx"
        lonely.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        legend = workbook.active
        legend.title = "Haeufigkeitsklassen"
        legend["A1"] = "Klassen"
        legend["B2"], legend["B3"] = "FanPower", "W"
        legend["A4"], legend["B4"] = 1, 10
        sheet = workbook.create_sheet("Daten Solo")
        sheet.cell(row=1, column=2, value="Case A")
        sheet.cell(row=1, column=3, value="Testgroessen")
        sheet.cell(row=1, column=4, value="Diag")
        sheet.cell(row=3, column=3, value="FanPower")
        sheet.cell(row=4, column=3, value=5)
        workbook.save(lonely)
        bands = build_split_header_distribution_bands(
            lonely,
            case_ids=["Case A"],
            quantities=[DistributionQuantity(header_label="FanPower", legend_key="FanPower")],
            scored_section="Testgroessen",
            diagnostic_section="Diag",
            candidate_sheet="Daten Testprogramm",
            data_prefix="Daten ",
        )
        self.assertEqual(bands, {})


@unittest.skipUnless(
    (REPO_ROOT / "SIA_4010_geteilter_Link" / "Test5").is_dir() and RUN_HEAVY,
    "set SIA4010_RUN_HEAVY=1 (with the official package) to run the heavy real extraction",
)
class Test5ReferenceBandTests(unittest.TestCase):
    """Build the real Test 5 ventilation scatter bands from the reference sheets."""

    @classmethod
    def setUpClass(cls):
        """Extract the Test 5 reference bands once."""
        import glob

        from swiss_sia.reference_model.sia4010.distribution_reference import (
            build_split_header_distribution_bands,
        )

        warnings.filterwarnings("ignore")
        entry = DISTRIBUTION_CRITERIA["5"]
        workbook = glob.glob(
            str(REPO_ROOT / "SIA_4010_geteilter_Link" / "Test5" / "Resultaterfassung*")
        )[0]
        cls.bands = build_split_header_distribution_bands(
            workbook,
            entry["case_ids"],
            entry["quantities"],
            scored_section=entry["scored_section"],
            diagnostic_section=entry["diagnostic_section"],
            candidate_sheet=entry["candidate_sheet"],
            data_prefix=entry["data_prefix"],
        )

    def test_scored_quantity_count_per_case_matches_the_workbook(self):
        counts = {}
        for case_id, _quantity in self.bands:
            counts[case_id] = counts.get(case_id, 0) + 1
        self.assertEqual(
            counts,
            {"Test 5A": 3, "Test 5B": 4, "Test 5C": 8, "Test 5D": 1},
        )

    def test_no_band_rests_on_a_single_program(self):
        for key, band in self.bands.items():
            with self.subTest(case=key[0], quantity=key[1]):
                self.assertGreaterEqual(band.program_count, 2)

    def test_diagnostic_fan_column_is_not_scored_for_case_5b(self):
        # Only the Daten Tas sheet omits the diagnostic boundary for 5B; the
        # agreed scored set must still exclude the fan power there.
        self.assertNotIn(
            ("Test 5B", "Leistung Zu- und Abluftventilator"), self.bands
        )


if __name__ == "__main__":
    unittest.main()
