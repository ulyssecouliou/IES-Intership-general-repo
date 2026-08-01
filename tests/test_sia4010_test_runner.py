"""End-to-end tests for the SIA 4010 test runner (bundle -> parse -> compare)."""

import unittest
from pathlib import Path

from openpyxl import Workbook

from swiss_sia.reference_model.exceptions import ConfigurationError
from swiss_sia.reference_model.sia4010.bundle_builder import write_manifest
from swiss_sia.reference_model.sia4010.compliance_comparator import (
    ComparisonOutcome,
    ComparisonStatus,
)
from swiss_sia.reference_model.sia4010.expected_results import ObservedResult
from swiss_sia.reference_model.sia4010.workbook_loaders import (
    parse_test1_reference_bands,
)
from swiss_sia.reference_model.sia4010.frequency_distribution import DistributionOutcome
from swiss_sia.reference_model.sia4010.test_runner import (
    Sia4010TestRunner,
    to_test_results_map,
)


def _dist(status: str, case_id: str = "") -> DistributionOutcome:
    """Build a distribution outcome fixture with a chosen status."""

    return DistributionOutcome(
        quantity="Q", unit="W", status=status, message="", case_id=case_id,
        out_of_band_bins=(), program_count=2, source_locator="loc",
    )

TEST_ROOT = Path(__file__).resolve().parents[1]
TEST_OUTPUT_ROOT = TEST_ROOT / ".codex_tmp" / "sia4010_runner_tests"


def _build_bundle(name: str) -> Path:
    """Create a checksum-verified bundle with a Test 1 evaluation workbook."""

    root = TEST_OUTPUT_ROOT / name
    (root / "Test1").mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung Testfälle"
    ws["A1"] = "Table 28 — Test results sensible energy"
    ws["A3"], ws["B3"], ws["D3"] = "Case id.", "1E", "kWh"
    for col, label in {
        "A": "Month", "B": "Testprogramm", "C": "IDA ICE", "G": "Mittelwert",
        "H": "obere Grenze", "I": "untere Grenze",
    }.items():
        ws["{}5".format(col)] = label
    ws["A6"], ws["G6"], ws["H6"], ws["I6"] = 1, 530.6, 588.9, 472.3
    ws["A7"], ws["G7"], ws["H7"], ws["I7"] = 2, 404.2, 452.7, 355.7
    wb.save(root / "Test1" / "Resultaterfassung_Test1.xlsx")
    write_manifest(root)
    return root


def _observed_from(root: Path, value_fn):
    """Build observed results keyed exactly to the parsed reference metrics."""

    workbook = root / "Test1" / "Resultaterfassung_Test1.xlsx"
    observed = []
    for expected in parse_test1_reference_bands(workbook):
        observed.append(
            ObservedResult(
                test_id=expected.test_id,
                case_id=expected.case_id,
                metric=expected.metric,
                value=value_fn(expected),
                unit=expected.unit,
                evidence_locator="VE APS extract",
            )
        )
    return observed


def _build_multi_bundle(name: str) -> Path:
    """Create a bundle with Test 1 and Test 2 evaluation workbooks."""

    root = TEST_OUTPUT_ROOT / name
    (root / "Test1").mkdir(parents=True, exist_ok=True)
    (root / "Test2").mkdir(parents=True, exist_ok=True)
    # Test 1 (case-per-block)
    wb1 = Workbook()
    ws1 = wb1.active
    ws1.title = "Zusammenfassung Testfälle"
    ws1["A1"] = "Table 28 — Test results sensible energy"
    ws1["A3"], ws1["B3"], ws1["D3"] = "Case id.", "1E", "kWh"
    for col, label in {
        "A": "Month", "B": "Testprogramm", "G": "Mittelwert",
        "H": "obere Grenze", "I": "untere Grenze",
    }.items():
        ws1["{}5".format(col)] = label
    ws1["A6"], ws1["G6"], ws1["H6"], ws1["I6"] = 1, 530.6, 588.9, 472.3
    wb1.save(root / "Test1" / "Resultaterfassung_Test1.xlsx")
    # Test 2 (case-per-row)
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "Zusammenfassung"
    for col, label in {
        "D": "Testprogramm", "M": "Mittelwert", "N": "obere Grenze", "O": "untere Grenze",
    }.items():
        ws2["{}9".format(col)] = label
    ws2["A12"], ws2["D12"] = "Fall", "Jahresenergie solar"
    ws2["A13"], ws2["N13"] = "Testfaelle", "kWh"
    ws2["A14"], ws2["M14"], ws2["N14"], ws2["O14"] = "Test 2 A", 1048.9, 1202.3, 895.5
    wb2.save(root / "Test2" / "Resultaterfassung_Test2.xlsx")
    write_manifest(root)
    return root


class Sia4010ClassBridgeTests(unittest.TestCase):
    def test_evaluate_all_and_shape_class_matrix_input(self):
        runner = Sia4010TestRunner()
        root = _build_multi_bundle("multi")
        evaluations = runner.evaluate_all(root, {})  # no observed -> NOT_CHECKABLE
        self.assertEqual(set(evaluations), {"1", "2"})
        class_input = to_test_results_map(evaluations)
        self.assertIn("test_1", class_input)
        self.assertIn("test_2", class_input)
        # With no observed results every band is NOT_CHECKABLE (never a false pass).
        self.assertEqual(class_input["test_1"]["status"], "NOT_CHECKABLE")
        self.assertEqual(class_input["test_2"]["status"], "NOT_CHECKABLE")

    def test_test2_exports_exact_variant_and_requires_distribution(self):
        runner = Sia4010TestRunner()
        root = _build_multi_bundle("multi_exact_variant")
        bands = runner.expected_bands(root, "2")
        observed = [
            ObservedResult(
                test_id=band.test_id,
                case_id=band.case_id,
                metric=band.metric,
                value=(band.lower_bound + band.upper_bound) / 2.0,
                unit=band.unit,
                evidence_locator="VE APS extract",
            )
            for band in bands
        ]
        without_distribution = runner.evaluate_test(root, "2", observed)
        self.assertEqual(
            to_test_results_map({"2": without_distribution})["test_2A"]["status"],
            "NOT_CHECKABLE",
        )

        complete = runner.evaluate_test(
            root,
            "2",
            observed,
            distribution_outcomes=[_dist("PASS", case_id="2A")],
        )
        shaped = to_test_results_map({"2": complete})
        self.assertEqual(shaped["test_2A"]["status"], "OFFICIAL_RESULTS_RECORDED")
        self.assertNotIn("test_2B", shaped)

    def test_exact_variant_keys_cover_tests_3_and_5(self):
        keys = Sia4010TestRunner._variant_key
        self.assertEqual(keys("3", "Test 3 A"), "test_3A")
        self.assertEqual(keys("3", "Test 3 L"), "test_3L")
        self.assertEqual(keys("5", "Test 5A"), "test_5A")
        self.assertEqual(keys("5", "Test 5D"), "test_5D")
        self.assertEqual(keys("4", "Test 4"), "test_4")

    def test_tests_3_and_5_cannot_pass_without_hourly_distribution(self):
        def passing(test_id, case_id):
            return ComparisonOutcome(
                key="{}|{}|annual".format(test_id, case_id),
                status=ComparisonStatus.PASS,
                message="",
                expected_value=1.0,
                observed_value=1.0,
                unit="kWh",
                absolute_difference=0.0,
                allowed_absolute_difference=None,
                source_locator="official workbook",
            )

        for test_id, case_id, variant in (
            ("3", "Test 3 A", "test_3A"),
            ("5", "Test 5D", "test_5D"),
        ):
            with self.subTest(test_id=test_id):
                statuses, bands, _counts = Sia4010TestRunner._variant_outcomes(
                    test_id, (passing(test_id, case_id),), ()
                )
                self.assertEqual(bands[variant], "OFFICIAL_RESULTS_RECORDED")
                self.assertEqual(statuses[variant], "NOT_CHECKABLE")


class Sia4010TestRunnerTests(unittest.TestCase):
    def setUp(self):
        self.runner = Sia4010TestRunner()

    def test_within_band_records_official_results_conservatively(self):
        root = _build_bundle("within")
        observed = _observed_from(root, lambda e: (e.lower_bound + e.upper_bound) / 2)
        evaluation = self.runner.evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "OFFICIAL_RESULTS_RECORDED")
        self.assertGreater(evaluation.counts["PASS"], 0)
        self.assertEqual(evaluation.counts["FAIL"], 0)

    def test_value_outside_band_fails(self):
        root = _build_bundle("outside")
        observed = _observed_from(root, lambda e: e.upper_bound + 100.0)
        evaluation = self.runner.evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "FAILED")
        self.assertGreater(evaluation.counts["FAIL"], 0)

    def test_no_observed_results_is_not_checkable(self):
        root = _build_bundle("missing")
        evaluation = self.runner.evaluate_test(root, "1", [])
        self.assertEqual(evaluation.status, "NOT_CHECKABLE")
        self.assertEqual(evaluation.counts["PASS"], 0)
        self.assertEqual(evaluation.counts["FAIL"], 0)

    def test_unregistered_test_id_is_rejected(self):
        root = _build_bundle("unregistered")
        with self.assertRaises(ConfigurationError):
            self.runner.evaluate_test(root, "9", [])

    def test_run_is_an_accessible_class_method(self):
        # Regression guard: `run` must be a real method, not mis-indented dead code.
        self.assertTrue(callable(getattr(Sia4010TestRunner, "run", None)))
        root = _build_bundle("legacy_run")
        result = self.runner.run(root, [])
        # No expected_results_csv role in this bundle -> empty comparisons, no crash.
        self.assertEqual(result.comparisons, ())
        self.assertIn("SIA", result.bundle.issued_by)


class Sia4010DistributionCombinationTests(unittest.TestCase):
    """Fold the frequency-distribution criterion into the test verdict, conservatively."""

    def test_combine_status_rules(self):
        combine = Sia4010TestRunner._combine_status
        self.assertEqual(
            combine("OFFICIAL_RESULTS_RECORDED", (_dist("PASS"),)),
            "OFFICIAL_RESULTS_RECORDED",
        )
        self.assertEqual(combine("OFFICIAL_RESULTS_RECORDED", (_dist("FAIL"),)), "FAILED")
        self.assertEqual(
            combine("OFFICIAL_RESULTS_RECORDED", (_dist("NOT_CHECKABLE"),)), "NOT_CHECKABLE"
        )
        self.assertEqual(combine("FAILED", (_dist("PASS"),)), "FAILED")
        self.assertEqual(combine("OFFICIAL_RESULTS_RECORDED", ()), "OFFICIAL_RESULTS_RECORDED")
        self.assertEqual(combine("NO_REFERENCE_BANDS", ()), "NO_REFERENCE_BANDS")

    def test_evaluate_test_folds_distribution_into_status(self):
        root = _build_bundle("dist")
        observed = _observed_from(root, lambda e: (e.lower_bound + e.upper_bound) / 2)
        runner = Sia4010TestRunner()
        # Band inside its band (OFFICIAL) but a not-checkable distribution must
        # pull the overall verdict down to NOT_CHECKABLE (never a false pass).
        evaluation = runner.evaluate_test(
            root, "1", observed, distribution_outcomes=[_dist("NOT_CHECKABLE")]
        )
        self.assertEqual(evaluation.band_status, "OFFICIAL_RESULTS_RECORDED")
        self.assertEqual(evaluation.status, "NOT_CHECKABLE")
        results = to_test_results_map({"1": evaluation})
        self.assertEqual(results["test_1"]["status"], "NOT_CHECKABLE")
        self.assertEqual(results["test_1"]["band_status"], "OFFICIAL_RESULTS_RECORDED")
        self.assertEqual(results["test_1"]["distribution"]["criteria"], 1)

    def test_distribution_pass_keeps_official_results_recorded(self):
        root = _build_bundle("dist_pass")
        observed = _observed_from(root, lambda e: (e.lower_bound + e.upper_bound) / 2)
        evaluation = Sia4010TestRunner().evaluate_test(
            root, "1", observed, distribution_outcomes=[_dist("PASS")]
        )
        self.assertEqual(evaluation.status, "OFFICIAL_RESULTS_RECORDED")

    def test_test2_accepts_one_annual_quantity_but_not_a_failed_quantity(self):
        def comparison(metric, status):
            return ComparisonOutcome(
                key="2|Test 2 A|{}".format(metric),
                status=status,
                message="",
                expected_value=1.0,
                observed_value=1.0 if status != ComparisonStatus.NOT_CHECKABLE else None,
                unit="kWh",
                absolute_difference=None,
                allowed_absolute_difference=None,
                source_locator="official workbook",
            )

        pass_one = (
            comparison("solar gain", ComparisonStatus.PASS),
            comparison("transmitted radiation", ComparisonStatus.NOT_CHECKABLE),
        )
        statuses, band_statuses, _ = Sia4010TestRunner._variant_outcomes(
            "2", pass_one, (_dist("PASS", case_id="2A"),)
        )
        self.assertEqual(
            band_statuses["test_2A"], "OFFICIAL_RESULTS_RECORDED"
        )
        self.assertEqual(statuses["test_2A"], "OFFICIAL_RESULTS_RECORDED")

        with_failure = pass_one + (
            comparison("supplied out-of-band value", ComparisonStatus.FAIL),
        )
        statuses, band_statuses, _ = Sia4010TestRunner._variant_outcomes(
            "2", with_failure, (_dist("PASS", case_id="2A"),)
        )
        self.assertEqual(band_statuses["test_2A"], "FAILED")
        self.assertEqual(statuses["test_2A"], "FAILED")


if __name__ == "__main__":
    unittest.main()
