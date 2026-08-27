"""Tests for the non-scored official Test 2 diagnostic 2E1 comparison."""

import hashlib
import json
import shutil
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from swiss_sia.reference_model.sia4010.test2a_diagnostic_evaluation import (
    NOT_CHECKABLE_STATUS,
    REFERENCE_ONLY_STATUS,
    TOTAL_GAIN_SERIES_ID,
    evaluate_test2a_2e1_diagnostic,
    load_test2a_2e1_reference_dataset,
    write_test2a_2e1_diagnostic_evaluation,
)
from swiss_sia.reference_model.sia4010.test2a_diagnostic_workbook import (
    load_test2a_diagnostic_workbook_binding,
)

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "SIA_4010_geteilter_Link" / "Test2" / "Resultaterfassung_Test2.xlsx"
WORK_ROOT = ROOT / ".codex_tmp" / "test2a_2e1_evaluation"


def _reference_program_series():
    binding = load_test2a_diagnostic_workbook_binding(WORKBOOK)
    workbook = load_workbook(WORKBOOK, data_only=True, read_only=True)
    try:
        worksheet = workbook["Daten_IDA_ICE Fe det Spec"]
        columns = {
            item.series_id: column_index_from_string(item.reference_column)
            for item in binding.series
        }
        first_column = min(columns.values())
        last_column = max(columns.values())
        result = {series_id: [] for series_id in columns}
        for row in worksheet.iter_rows(
            min_row=4,
            max_row=8763,
            min_col=first_column,
            max_col=last_column,
            values_only=True,
        ):
            for series_id, column in columns.items():
                result[series_id].append(row[column - first_column])
        return {series_id: tuple(values) for series_id, values in result.items()}
    finally:
        workbook.close()


@unittest.skipUnless(WORKBOOK.is_file(), "official Test 2 workbook unavailable")
class Test2A2E1DiagnosticEvaluationTests(unittest.TestCase):
    """Reference alignment is useful but can never become a compliance PASS."""

    @classmethod
    def setUpClass(cls):
        cls.reference = load_test2a_2e1_reference_dataset(WORKBOOK)
        cls.series = _reference_program_series()
        cls.required_ids = tuple(cls.series)

    def setUp(self):
        self.output_dir = (
            WORK_ROOT
            / hashlib.sha256(self._testMethodName.encode("utf-8")).hexdigest()[:12]
        )
        if self.output_dir.exists():
            shutil.rmtree(self.output_dir)
        self.output_dir.mkdir(parents=True)

    def tearDown(self):
        if self.output_dir.exists():
            shutil.rmtree(self.output_dir)

    def test_loads_the_plotted_annual_values_and_histogram_scatter(self):
        total = self.reference.annual(TOTAL_GAIN_SERIES_ID)
        self.assertEqual(len(total.program_values), 6)
        self.assertAlmostEqual(total.minimum, 315.61987999999894)
        self.assertAlmostEqual(total.maximum, 543.7735296599999)
        self.assertIsNone(total.acceptance_criterion)
        self.assertEqual(
            self.reference.total_gain_distribution.program_count,
            6,
        )
        self.assertFalse(self.reference.acceptance_criterion_available)

    def test_reference_program_is_aligned_but_never_a_compliance_pass(self):
        evaluation = evaluate_test2a_2e1_diagnostic(
            self.series,
            self.reference,
            self.required_ids,
        )
        self.assertEqual(evaluation.status, REFERENCE_ONLY_STATUS)
        self.assertEqual(
            evaluation.technical_alignment,
            "WITHIN_TECHNICAL_REFERENCE_ENVELOPE",
        )
        self.assertTrue(evaluation.required_eight_series_complete)
        self.assertTrue(evaluation.engineering_review_ready)
        self.assertFalse(evaluation.acceptance_criterion_available)
        self.assertFalse(evaluation.compliance_pass)
        self.assertFalse(evaluation.optical_mapping_qualified)
        self.assertFalse(evaluation.dynamic_control_qualified)

    def test_zeroed_candidate_is_outside_but_still_not_a_sia_fail(self):
        zeros = {series_id: (0.0,) * 8760 for series_id in self.required_ids}
        evaluation = evaluate_test2a_2e1_diagnostic(
            zeros,
            self.reference,
            self.required_ids,
        )
        self.assertEqual(evaluation.status, REFERENCE_ONLY_STATUS)
        self.assertEqual(
            evaluation.technical_alignment,
            "OUTSIDE_TECHNICAL_REFERENCE_ENVELOPE",
        )
        self.assertFalse(evaluation.engineering_review_ready)
        self.assertFalse(evaluation.compliance_pass)

    def test_missing_comparable_output_remains_not_checkable(self):
        incident_only = {
            series_id: values
            for series_id, values in self.series.items()
            if series_id.startswith("hourly_incident_")
        }
        evaluation = evaluate_test2a_2e1_diagnostic(
            incident_only,
            self.reference,
            self.required_ids,
        )
        self.assertEqual(evaluation.status, NOT_CHECKABLE_STATUS)
        self.assertEqual(evaluation.technical_alignment, "NOT_CHECKABLE")
        self.assertFalse(evaluation.engineering_review_ready)

    def test_writes_a_checksummed_audit_without_validation_claim(self):
        evaluation = evaluate_test2a_2e1_diagnostic(
            self.series,
            self.reference,
            self.required_ids,
        )
        path = self.output_dir / "evaluation.json"
        write_test2a_2e1_diagnostic_evaluation(
            path,
            evaluation,
            self.reference,
        )
        payload = json.loads(path.read_text(encoding="utf-8"))
        self.assertFalse(payload["compliance_pass"])
        self.assertFalse(payload["acceptance_criterion_available"])
        checksum_path = path.with_suffix(".json.sha256")
        expected = hashlib.sha256(path.read_bytes()).hexdigest()
        self.assertTrue(checksum_path.read_text(encoding="ascii").startswith(expected))


if __name__ == "__main__":
    unittest.main()
