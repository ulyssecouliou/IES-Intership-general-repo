"""Pure-Python dry run of the full SIA 4010 Phase B pipeline.

Fake VE source -> observed results -> runner -> verdict. This proves the
ingestion/comparison chain end to end and that observed keys align with the
official expected keys, without a real VE runtime.
"""

import unittest
from pathlib import Path

from openpyxl import Workbook

from swiss_sia.reference_model.sia4010.bundle_builder import write_manifest
from swiss_sia.reference_model.sia4010.expected_results import ExpectedResult
from swiss_sia.reference_model.sia4010.observed_extraction import (
    DictResultSource,
    MetricBinding,
    VeApsResolver,
    VeApsResultAccessor,
    build_observed_results,
)
from swiss_sia.reference_model.sia4010.test_runner import Sia4010TestRunner
from swiss_sia.reference_model.sia4010.workbook_loaders import (
    deduplicate_expected_keys,
    parse_test1_reference_bands,
)


def _band_expected(metric="M", lower=90.0, upper=110.0):
    return ExpectedResult(
        test_id="1", case_id="1E", metric=metric, expected_value=100.0, unit="kWh",
        absolute_tolerance=None, relative_tolerance=None,
        source_locator="loc", source_checksum="chk",
        lower_bound=lower, upper_bound=upper,
    )

TEST_ROOT = Path(__file__).resolve().parents[1]
TEST_OUTPUT_ROOT = TEST_ROOT / ".codex_tmp" / "sia4010_phase_b_tests"


def _build_bundle(name: str) -> Path:
    """Create a Test 1 bundle with a case-1E band block."""

    root = TEST_OUTPUT_ROOT / name
    (root / "Test1").mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Zusammenfassung Testfälle"
    ws["A1"] = "Table 28 — Test results sensible energy"
    ws["A3"], ws["B3"], ws["D3"] = "Case id.", "1E", "kWh"
    for col, label in {
        "A": "Month", "B": "Testprogramm", "G": "Mittelwert",
        "H": "obere Grenze", "I": "untere Grenze",
    }.items():
        ws["{}5".format(col)] = label
    ws["A6"], ws["G6"], ws["H6"], ws["I6"] = 1, 530.6, 588.9, 472.3
    ws["A7"], ws["G7"], ws["H7"], ws["I7"] = 2, 404.2, 452.7, 355.7
    wb.save(root / "Test1" / "Resultaterfassung_Test1.xlsx")
    write_manifest(root)
    return root


def _expected(root: Path):
    return parse_test1_reference_bands(root / "Test1" / "Resultaterfassung_Test1.xlsx")


class Sia4010PhaseBDryRunTests(unittest.TestCase):
    def setUp(self):
        self.runner = Sia4010TestRunner()

    def test_within_band_dry_run_records_results(self):
        root = _build_bundle("within")
        expected = _expected(root)
        source = DictResultSource(
            {e.key: ((e.lower_bound + e.upper_bound) / 2, e.unit) for e in expected}
        )
        observed = build_observed_results(expected, source)
        self.assertEqual(len(observed), len(expected))  # keys align one-to-one
        evaluation = self.runner.evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "OFFICIAL_RESULTS_RECORDED")
        self.assertEqual(evaluation.counts["FAIL"], 0)

    def test_out_of_band_dry_run_fails(self):
        root = _build_bundle("outof")
        expected = _expected(root)
        source = DictResultSource(
            {e.key: (e.upper_bound + 1000.0, e.unit) for e in expected}
        )
        observed = build_observed_results(expected, source)
        evaluation = self.runner.evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "FAILED")

    def test_missing_ve_value_stays_not_checkable_without_fabrication(self):
        root = _build_bundle("partial")
        expected = _expected(root)
        # Resolver supplies only the first metric; the rest are omitted.
        first_key = expected[0].key
        source = DictResultSource(
            {first_key: (
                (expected[0].lower_bound + expected[0].upper_bound) / 2,
                expected[0].unit,
            )}
        )
        observed = build_observed_results(expected, source)
        self.assertEqual(len(observed), 1)  # nothing fabricated for the others
        evaluation = self.runner.evaluate_test(root, "1", observed)
        self.assertGreater(evaluation.counts["NOT_CHECKABLE"], 0)


class VeApsResolverTests(unittest.TestCase):
    def test_unconfirmed_binding_resolves_to_none(self):
        binding = MetricBinding(quantity="q", unit="kWh", confirmed=False)
        resolver = VeApsResolver(lambda b, e: 100.0, {("1", "M"): binding})
        self.assertIsNone(resolver(_band_expected()))

    def test_missing_binding_resolves_to_none(self):
        resolver = VeApsResolver(lambda b, e: 100.0, {})
        self.assertIsNone(resolver(_band_expected()))

    def test_confirmed_binding_but_absent_value_resolves_to_none(self):
        binding = MetricBinding(quantity="q", unit="kWh", confirmed=True)
        resolver = VeApsResolver(lambda b, e: None, {("1", "M"): binding})
        self.assertIsNone(resolver(_band_expected()))

    def test_confirmed_binding_resolves_value_and_unit(self):
        binding = MetricBinding(quantity="q", unit="kWh", confirmed=True)
        resolver = VeApsResolver(lambda b, e: 101.5, {("1", "M"): binding})
        self.assertEqual(resolver(_band_expected()), (101.5, "kWh"))

    def test_end_to_end_pipeline_through_resolver(self):
        root = _build_bundle("resolver")
        expected = deduplicate_expected_keys(_expected(root))
        bindings = {
            ("1", e.metric): MetricBinding(quantity=e.metric, unit=e.unit, confirmed=True)
            for e in expected
        }
        mid = {e.key: (e.lower_bound + e.upper_bound) / 2 for e in expected}
        resolver = VeApsResolver(lambda b, e: mid.get(e.key), bindings)
        observed = build_observed_results(expected, resolver)
        evaluation = Sia4010TestRunner().evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "OFFICIAL_RESULTS_RECORDED")


class VeApsResultAccessorTests(unittest.TestCase):
    def test_unregistered_quantity_returns_none(self):
        accessor = VeApsResultAccessor(aps_results=object(), quantity_extractors={})
        binding = MetricBinding(quantity="fan_energy", unit="kWh", confirmed=True)
        self.assertIsNone(accessor(binding, _band_expected()))

    def test_registered_extractor_returns_value(self):
        accessor = VeApsResultAccessor(aps_results={"fan_energy": 100.0})
        accessor.register("fan_energy", lambda aps, b, e: aps.get(b.quantity))
        binding = MetricBinding(quantity="fan_energy", unit="kWh", confirmed=True)
        self.assertEqual(accessor(binding, _band_expected()), 100.0)

    def test_extractor_error_is_fail_closed(self):
        def boom(aps, binding, expected):
            raise RuntimeError("runtime read failed")

        accessor = VeApsResultAccessor(aps_results=None, quantity_extractors={"q": boom})
        self.assertIsNone(
            accessor(MetricBinding(quantity="q", unit="kWh", confirmed=True), _band_expected())
        )

    def test_end_to_end_via_resolver_and_accessor(self):
        root = _build_bundle("accessor")
        expected = deduplicate_expected_keys(_expected(root))
        mid = {e.key: (e.lower_bound + e.upper_bound) / 2 for e in expected}
        accessor = VeApsResultAccessor(aps_results=mid)
        accessor.register("band_mid", lambda aps, b, e: aps.get(e.key))
        bindings = {
            ("1", e.metric): MetricBinding(quantity="band_mid", unit=e.unit, confirmed=True)
            for e in expected
        }
        resolver = VeApsResolver(accessor, bindings)
        observed = build_observed_results(expected, resolver)
        evaluation = Sia4010TestRunner().evaluate_test(root, "1", observed)
        self.assertEqual(evaluation.status, "OFFICIAL_RESULTS_RECORDED")


if __name__ == "__main__":
    unittest.main()
