import openpyxl
import json
import os
from pathlib import Path

REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
SIA_SOURCE_ROOT = Path(
    os.environ.get("SIA_4010_DOSSIER", REPOSITORY_ROOT / "SIA_4010_geteilter_Link")
)
wb_path = SIA_SOURCE_ROOT / "Test1" / "Resultaterfassung_Test1.xlsx"
wb = openpyxl.load_workbook(wb_path, data_only=True)
ws = wb["Zusammenfassung Testfälle"]

# ===== TABLE 30 EXTRACTION =====
table_30_data = {}


def get_month_name(month_num):
    """Convert 1-12 to month_01 to month_12"""
    return f"month_{month_num:02d}"


# Blocks FF (monthly with Month column)
cases_30 = [
    ("600FF", "AG", "AH", "AI", "AJ", "AK", "AL", "AM"),
    ("900FF", "AP", "AQ", "AR", "AS", "AT", "AU", "AV"),
]

for (
    case_id,
    col_month,
    col_testprog,
    col_iso,
    col_ida,
    col_excel,
    col_energyplus,
    col_tas,
) in cases_30:
    case_data = {
        "_metadata": {
            "case_id": case_id,
            "table_number": 30,
            "description": "Free-floating temperature (monthly)",
            "columns": {
                "month": col_month,
                "testprogramm": col_testprog,
                "iso_52016_1_2017": col_iso,
                "ida_ice_5_0_beta_23": col_ida,
                "excel_sia_380_2": col_excel,
                "energyplus_openstudio": col_energyplus,
                "tas_edsl": col_tas,
            },
        },
        "monthly": {},
    }

    # Rows 58-69 (Jan-Dec)
    for row in range(58, 70):
        month_num = ws[f"{col_month}{row}"].value
        testprog = ws[f"{col_testprog}{row}"].value
        iso = ws[f"{col_iso}{row}"].value
        ida = ws[f"{col_ida}{row}"].value
        excel = ws[f"{col_excel}{row}"].value
        energyplus = ws[f"{col_energyplus}{row}"].value
        tas = ws[f"{col_tas}{row}"].value

        if month_num and isinstance(month_num, int) and 1 <= month_num <= 12:
            month_key = get_month_name(month_num)
            case_data["monthly"][month_key] = {
                "testprogramm_candidate": {
                    "value": testprog if not isinstance(testprog, str) else None,
                    "unit": "C",
                    "cell": f"{col_testprog}{row}",
                    "note": "Excel error: #DIV/0!" if testprog == "#DIV/0!" else None,
                },
                "iso_52016_1_2017_reference": {
                    "value": iso,
                    "unit": "C",
                    "cell": f"{col_iso}{row}",
                },
                "ida_ice_5_0_beta_23": {
                    "value": ida,
                    "unit": "C",
                    "cell": f"{col_ida}{row}",
                },
                "excel_sia_380_2": {
                    "value": excel,
                    "unit": "C",
                    "cell": f"{col_excel}{row}",
                },
                "energyplus_openstudio": {
                    "value": energyplus,
                    "unit": "C",
                    "cell": f"{col_energyplus}{row}",
                },
                "tas_edsl": {
                    "value": tas,
                    "unit": "C",
                    "cell": f"{col_tas}{row}",
                },
            }

    # Row 70 (Annual)
    row = 70
    month_num = ws[f"{col_month}{row}"].value
    testprog = ws[f"{col_testprog}{row}"].value
    iso = ws[f"{col_iso}{row}"].value
    ida = ws[f"{col_ida}{row}"].value
    excel = ws[f"{col_excel}{row}"].value
    energyplus = ws[f"{col_energyplus}{row}"].value
    tas = ws[f"{col_tas}{row}"].value

    if month_num == "Annual":
        case_data["monthly"]["annual"] = {
            "testprogramm_candidate": {
                "value": testprog if not isinstance(testprog, str) else None,
                "unit": "C",
                "cell": f"{col_testprog}{row}",
                "note": "Excel error: #DIV/0!" if testprog == "#DIV/0!" else None,
            },
            "iso_52016_1_2017_reference": {
                "value": iso,
                "unit": "C",
                "cell": f"{col_iso}{row}",
            },
            "ida_ice_5_0_beta_23": {
                "value": ida,
                "unit": "C",
                "cell": f"{col_ida}{row}",
            },
            "excel_sia_380_2": {
                "value": excel,
                "unit": "C",
                "cell": f"{col_excel}{row}",
            },
            "energyplus_openstudio": {
                "value": energyplus,
                "unit": "C",
                "cell": f"{col_energyplus}{row}",
            },
            "tas_edsl": {
                "value": tas,
                "unit": "C",
                "cell": f"{col_tas}{row}",
            },
        }

    table_30_data[case_id] = case_data

print("TABLE 30: 600FF/900FF monthly extracted with CORRECT columns")

# ===== TABLE 32 EXTRACTION =====
table_32_data = {}

cases_32 = [
    ("600FF", ("A", "B", "C", "D", "E", "F", "G")),
    ("900FF", ("I", "J", "K", "L", "M", "N", "O")),
]

for case_id, cols in cases_32:
    col_label, col_testprog, col_iso, col_ida, col_excel, col_energyplus, col_tas = cols

    case_data = {
        "_metadata": {
            "case_id": case_id,
            "table_number": 32,
            "description": "Annual extremes (Max/Min/Average)",
            "columns": {
                "label": col_label,
                "testprogramm": col_testprog,
                "iso_52016_1_2017": col_iso,
                "ida_ice_5_0_beta_23": col_ida,
                "excel_sia_380_2": col_excel,
                "energyplus_openstudio": col_energyplus,
                "tas_edsl": col_tas,
            },
        },
        "extremes": {},
    }

    # Rows 105 (Max), 106 (Min), 107 (Average)
    for row, label_key in [(105, "max"), (106, "min"), (107, "average")]:
        label = ws[f"{col_label}{row}"].value
        testprog = ws[f"{col_testprog}{row}"].value
        iso = ws[f"{col_iso}{row}"].value
        ida = ws[f"{col_ida}{row}"].value
        excel = ws[f"{col_excel}{row}"].value
        energyplus = ws[f"{col_energyplus}{row}"].value
        tas = ws[f"{col_tas}{row}"].value

        case_data["extremes"][label_key] = {
            "label": label,
            "testprogramm_candidate": {
                "value": testprog if not isinstance(testprog, str) else None,
                "unit": "C",
                "cell": f"{col_testprog}{row}",
                "note": "Excel error: #DIV/0!" if testprog == "#DIV/0!" else None,
            },
            "iso_52016_1_2017_reference": {
                "value": iso,
                "unit": "C",
                "cell": f"{col_iso}{row}",
            },
            "ida_ice_5_0_beta_23": {
                "value": ida,
                "unit": "C",
                "cell": f"{col_ida}{row}",
            },
            "excel_sia_380_2": {
                "value": excel,
                "unit": "C",
                "cell": f"{col_excel}{row}",
            },
            "energyplus_openstudio": {
                "value": energyplus,
                "unit": "C",
                "cell": f"{col_energyplus}{row}",
            },
            "tas_edsl": {
                "value": tas,
                "unit": "C",
                "cell": f"{col_tas}{row}",
            },
        }

    table_32_data[case_id] = case_data

print("TABLE 32: 600FF/900FF extremes extracted with CORRECT columns")

# Save to JSON files for inspection
output_dir = Path(__file__).resolve().parent
with (output_dir / "table_30_corrected.json").open("w", encoding="utf-8") as f:
    json.dump(table_30_data, f, indent=2)

with (output_dir / "table_32_corrected.json").open("w", encoding="utf-8") as f:
    json.dump(table_32_data, f, indent=2)

print(f"\nFiles saved to {output_dir}/")
print("  - table_30_corrected.json")
print("  - table_32_corrected.json")
