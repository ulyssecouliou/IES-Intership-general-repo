# -*- coding: utf-8 -*-
"""Extracts the Zürich-Kloten outdoor temperature from the OFFICIAL SIA workbooks.

DISCOVERY OF 2026-08-05. The official evaluation workbook for Test 4,
`Resultaterfassung Test4.xlsx`, contains a **`Wetterdaten`** sheet that
had gone unnoticed until now. It carries 8760 hourly values:

    column 1: "Site Outdoor Air Drybulb Temperature [C](Hourly)"
    column 2: "EMS Two Day Average OA Temp [C](Hourly)"

Column 1 is the **hourly outdoor air temperature for Zürich-Kloten** as
the reference program EnergyPlus read it from the original SIA 2028 file.
Column 2 is the **48-hour running mean**, already calculated — that is
exactly the x-axis of figure 1 of SIA 380/2.

SCOPE AND LIMITATIONS — read before using.

What this gives: a temperature series of official SIA origin, without
purchase. Sufficient for anything that depends only on air temperature —
in particular the dry cooler and the outdoor-air heat exchanger of Test 7.

What this does NOT give: solar radiation. The original SIA 2028 file
contains irradiance on vertical surfaces for the main orientations
(from the Test 2 EXCEL report); none of that is here. Tests 1E, 2, 3, 4,
5, 6 depend on it, and Test 7 depends on it for its mandatory quantity n° 14
"Elektrische Energie PV".

This is therefore NOT a substitute for the full SIA 2028 dataset. It is
an authentic and verifiable piece of the puzzle, from an official source.

Usage:
    python scripts/extract_kloten_weather_from_sia.py [--ecrire]
"""

from __future__ import print_function

import io
import json
import os
import sys

import openpyxl

_ICI = os.path.dirname(os.path.abspath(__file__))
_RACINE = os.path.abspath(os.path.join(_ICI, os.pardir))

# The official folder lives in the other repository (118 MB, excluded by .gitignore).
_DOSSIER_SIA = os.environ.get(
    "SIA_4010_DOSSIER", os.path.join(_RACINE, "SIA_4010_geteilter_Link")
)

_CLASSEUR = os.path.join(_DOSSIER_SIA, "Test4", "Resultaterfassung Test4.xlsx")
_FEUILLE = "Wetterdaten"

_SORTIE_JSON = os.path.join(
    _RACINE, "refs", "reference-data", "sia-2028-kloten-temperature.json"
)
_SORTIE_CSV = os.path.join(
    _RACINE, "refs", "reference-data", "sia-2028-kloten-temperature.csv"
)

HEURES = 8760

# Plausibility bounds for Zürich-Kloten (DRY reference year).
# They do not validate the exact value: they catch an extraction that
# picked the wrong column or the wrong sheet.
MOYENNE_ATTENDUE = (8.0, 11.0)
MIN_ATTENDU = (-25.0, -5.0)
MAX_ATTENDU = (28.0, 40.0)


class ExtractionRefusee(RuntimeError):
    pass


def extraire():
    if not os.path.exists(_CLASSEUR):
        raise ExtractionRefusee(
            "classeur officiel introuvable : %s\n"
            "Définir SIA_4010_DOSSIER si le dossier SIA est ailleurs." % _CLASSEUR
        )

    classeur = openpyxl.load_workbook(_CLASSEUR, data_only=True, read_only=True)
    if _FEUILLE not in classeur.sheetnames:
        raise ExtractionRefusee(
            "feuille %r absente de %s" % (_FEUILLE, os.path.basename(_CLASSEUR))
        )
    feuille = classeur[_FEUILLE]

    lignes = [list(r) for r in feuille.iter_rows(values_only=True)]
    classeur.close()

    entete = [str(v) for v in lignes[0][:2]]
    if "Drybulb" not in entete[0]:
        raise ExtractionRefusee(
            "colonne 1 inattendue : %r — l'extraction est refusée plutôt que "
            "de figer une grandeur non identifiée" % entete[0]
        )

    temperature, moyenne_48h = [], []
    for ligne in lignes[1:]:
        if isinstance(ligne[0], (int, float)):
            temperature.append(float(ligne[0]))
            moyenne_48h.append(
                float(ligne[1]) if isinstance(ligne[1], (int, float)) else None
            )

    if len(temperature) != HEURES:
        raise ExtractionRefusee(
            "%d heures extraites, %d attendues" % (len(temperature), HEURES)
        )

    moyenne = sum(temperature) / len(temperature)
    mini, maxi = min(temperature), max(temperature)
    for valeur, (bas, haut), nom in (
        (moyenne, MOYENNE_ATTENDUE, "moyenne"),
        (mini, MIN_ATTENDU, "minimum"),
        (maxi, MAX_ATTENDU, "maximum"),
    ):
        if not (bas <= valeur <= haut):
            raise ExtractionRefusee(
                "%s = %.2f °C hors de la plage de vraisemblance [%.1f ; %.1f] "
                "pour Zürich-Kloten" % (nom, valeur, bas, haut)
            )

    return entete, temperature, moyenne_48h


def construire(entete, temperature, moyenne_48h):
    moyenne = sum(temperature) / len(temperature)
    return {
        "grandeur": "température d'air extérieur horaire, Zürich-Kloten",
        "statut": "FIGÉ — source officielle SIA, extrait le 2026-08-05",
        "source": {
            "fichier": "SIA_4010_geteilter_Link/Test4/Resultaterfassung Test4.xlsx",
            "feuille": _FEUILLE,
            "colonnes": entete,
            "nature": "sortie du programme de référence EnergyPlus, telle que "
            "livrée par le SIA dans son classeur d'évaluation officiel",
            "origine_amont": "fichier SIA 2028 DRY normal, station Kloten "
            "(« Original-SIA-Datei » des rapports d'application)",
        },
        "heures": len(temperature),
        "agregats": {
            "min_c": round(min(temperature), 4),
            "max_c": round(max(temperature), 4),
            "moyenne_c": round(moyenne, 4),
        },
        "moyenne_glissante_48h": {
            "presente": any(v is not None for v in moyenne_48h),
            "colonne": entete[1] if len(entete) > 1 else None,
            "usage": "abscisse de la figure 1 de SIA 380/2 — permet de "
            "contrôler `engine.setpoint_curves.running_mean_48h` "
            "contre une série calculée par un programme de référence, "
            "y compris sur les 47 premières heures où la norme ne dit "
            "rien.",
        },
        "ce_que_cela_debloque": [
            "Tout ce qui ne dépend que de la température d'air extérieur : "
            "refroidisseur sec et échangeur sur air extérieur du Test 7, donc "
            "le point de fonctionnement de la pompe à chaleur.",
            "La validation de notre moyenne glissante sur 48 h.",
        ],
        "ce_que_cela_ne_debloque_pas": [
            "Le RAYONNEMENT SOLAIRE, absent de cette feuille. Le fichier "
            "SIA 2028 d'origine contient l'irradiance sur les surfaces "
            "verticales des orientations principales ; rien de cela n'est ici.",
            "Test 7 grandeur obligatoire n° 14 « Elektrische Energie PV » : "
            "exige l'irradiance sur le plan des modules.",
            "Tests 1E, 2, 3, 4, 5, 6 : tous solaires.",
            "L'humidité de l'air extérieur.",
        ],
        "avertissement": "Ce n'est PAS un substitut au jeu SIA 2028 complet. "
        "C'est une pièce authentique du puzzle, d'origine "
        "officielle, qui couvre la seule température.",
    }


def main():
    entete, temperature, moyenne_48h = extraire()
    donnees = construire(entete, temperature, moyenne_48h)

    print("feuille  : %s" % _FEUILLE)
    print("colonnes : %s" % " | ".join(entete))
    print("heures   : %d" % donnees["heures"])
    agregats = donnees["agregats"]
    print(
        "min %.2f   max %.2f   moyenne %.3f °C"
        % (agregats["min_c"], agregats["max_c"], agregats["moyenne_c"])
    )
    print(
        "moyenne glissante 48 h présente : %s"
        % donnees["moyenne_glissante_48h"]["presente"]
    )

    if "--ecrire" in sys.argv:
        with io.open(_SORTIE_JSON, "w", encoding="utf-8") as f:
            f.write(json.dumps(donnees, ensure_ascii=False, indent=2))
            f.write("\n")
        with io.open(_SORTIE_CSV, "w", encoding="utf-8") as f:
            f.write("heure,theta_e_air_c,moyenne_glissante_48h_c\n")
            for i, (t, m) in enumerate(zip(temperature, moyenne_48h), start=1):
                f.write("%d,%.6f,%s\n" % (i, t, "" if m is None else "%.6f" % m))
        print()
        print("écrit : %s" % os.path.relpath(_SORTIE_JSON, _RACINE))
        print("écrit : %s" % os.path.relpath(_SORTIE_CSV, _RACINE))


if __name__ == "__main__":
    main()
